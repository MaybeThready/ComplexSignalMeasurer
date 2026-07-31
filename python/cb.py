#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Vpp/频率校准分析工具。

兼容旧用法：
    python cb.py [输入文件] [输出目录]

推荐用法：
    python cb.py 校准.xlsx cb_output --cv frequency

输入列（单位固定）：
    测量值：mV Vpp
    真值：  mV Vpp
    频率：  kHz

相较旧版本，本程序有意不再为每条 Vpp 数据合成 2000 点正弦波。合成波形并没有
增加观测信息；对零点对称的奇函数校准，峰峰值可直接解析表示。当前模型均满足
Vmeas=0 -> Vcal=0，并使用归一化基函数和线性最小二乘稳定求解。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from rich.console import Console
from rich.panel import Panel
from rich.table import Table
from scipy.stats import t as student_t


APP_VERSION = "3.0"
RANDOM_SEED = 42
EPS = np.finfo(float).eps
console = Console(highlight=False)


# --------------------------- 字体与通用工具 ---------------------------
def configure_plot_font() -> bool:
    """配置中文字体；返回是否找到了可用中文字体。"""
    available = {item.name for item in font_manager.fontManager.ttflist}
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
        "SimSun",
    ]
    selected = next((name for name in candidates if name in available), None)
    if selected:
        plt.rcParams["font.sans-serif"] = [selected, "DejaVu Sans"]
    else:
        plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 120
    return selected is not None


HAS_CHINESE_FONT = configure_plot_font()


def plot_text(chinese: str, english: str) -> str:
    return chinese if HAS_CHINESE_FONT else english


def finite_float(value: Any) -> float | None:
    """把数值转换为 JSON 兼容的有限浮点数。"""
    if value is None:
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def to_builtin(value: Any) -> Any:
    """递归转换 numpy/pandas 对象，且不向 JSON 写入 NaN/Infinity。"""
    if isinstance(value, dict):
        return {str(key): to_builtin(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [to_builtin(item) for item in value]
    if isinstance(value, np.ndarray):
        return [to_builtin(item) for item in value.tolist()]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return finite_float(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    return value


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_percent(numerator: float, denominator: float) -> float | None:
    if abs(denominator) <= EPS:
        return None
    return 100.0 * numerator / denominator


def format_number(value: Any, digits: int = 4, missing: str = "—") -> str:
    number = finite_float(value)
    if number is None:
        return missing
    return f"{number:.{digits}f}"


def format_r2(value: Any) -> str:
    number = finite_float(value)
    return "—" if number is None else f"{number:.6f}"


# --------------------------- 数据加载与验证 ---------------------------
def normalize_header(value: Any) -> str:
    text = str(value).strip().lower()
    return re.sub(r"[\s_()（）\[\]【】,，/\\\-]+", "", text)


COLUMN_ALIASES = {
    "测量值": {
        "测量值",
        "测量值vppmv",
        "实测值",
        "实测值vppmv",
        "measurement",
        "measured",
        "measuredvpp",
        "measuredvppmv",
    },
    "真值": {
        "真值",
        "真值vppmv",
        "标准值",
        "标准值vppmv",
        "参考值",
        "reference",
        "referencevpp",
        "true",
        "truevpp",
        "truevppmv",
    },
    "频率": {
        "频率",
        "频率khz",
        "频率k赫兹",
        "freq",
        "freqkhz",
        "frequency",
        "frequencykhz",
    },
}


def find_required_columns(columns: Iterable[Any]) -> dict[str, Any]:
    normalized = {column: normalize_header(column) for column in columns}
    result: dict[str, Any] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        matches = [column for column, key in normalized.items() if key in aliases]
        if not matches:
            available = "、".join(str(column) for column in columns)
            raise ValueError(
                f"缺少“{canonical}”列。当前列为：{available}。"
                "频率必须以 kHz、幅值必须以 mV Vpp 表示。"
            )
        if len(matches) > 1:
            raise ValueError(
                f"有多列都可能是“{canonical}”："
                + "、".join(str(column) for column in matches)
            )
        result[canonical] = matches[0]
    return result


def read_tabular_file(path: Path, sheet: str | int = 0) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls", ".xlsm"}:
        try:
            return pd.read_excel(path, sheet_name=sheet)
        except ImportError as exc:
            raise RuntimeError(
                "读取 Excel 需要 openpyxl；请执行：pip install openpyxl"
            ) from exc
    if suffix in {".csv", ".txt", ".tsv"}:
        separator = "\t" if suffix == ".tsv" else None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return pd.read_csv(
                    path, sep=separator, engine="python", encoding=encoding
                )
            except UnicodeDecodeError:
                continue
        raise ValueError("文本文件编码无法识别，请保存为 UTF-8 或 GB18030。")
    raise ValueError("仅支持 .xlsx/.xls/.xlsm/.csv/.tsv/.txt 数据文件。")


def load_data(
    file_path: str | Path, sheet: str | int = 0
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """读取并严格验证校准数据，返回规范化数据和质量审计信息。"""
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"输入文件不存在：{path}")

    raw = read_tabular_file(path, sheet)
    if raw.empty:
        raise ValueError("输入文件没有数据。")

    mapping = find_required_columns(raw.columns)
    selected = raw[[mapping["测量值"], mapping["真值"], mapping["频率"]]].copy()
    selected.columns = ["测量值", "真值", "频率"]
    selected["源行号"] = np.arange(2, len(selected) + 2)

    fully_blank = selected[["测量值", "真值", "频率"]].isna().all(axis=1)
    blank_rows_removed = int(fully_blank.sum())
    selected = selected.loc[~fully_blank].copy()
    if selected.empty:
        raise ValueError("必需列中没有有效数据。")

    for column in ["测量值", "真值", "频率"]:
        selected[column] = pd.to_numeric(selected[column], errors="coerce")

    values = selected[["测量值", "真值", "频率"]].to_numpy(dtype=float)
    invalid = selected[["测量值", "真值", "频率"]].isna().any(axis=1)
    invalid |= ~np.isfinite(values).all(axis=1)
    if invalid.any():
        rows = "、".join(str(int(row)) for row in selected.loc[invalid, "源行号"][:10])
        suffix = "……" if int(invalid.sum()) > 10 else ""
        raise ValueError(f"第 {rows}{suffix} 行含空值、非数值或无穷值。")

    if (selected["测量值"] < 0).any():
        rows = selected.loc[selected["测量值"] < 0, "源行号"].tolist()[:10]
        raise ValueError(f"测量值 Vpp 不得为负；异常源行：{rows}")
    if (selected["真值"] < 0).any():
        rows = selected.loc[selected["真值"] < 0, "源行号"].tolist()[:10]
        raise ValueError(f"真值 Vpp 不得为负；异常源行：{rows}")
    if (selected["频率"] <= 0).any():
        rows = selected.loc[selected["频率"] <= 0, "源行号"].tolist()[:10]
        raise ValueError(f"频率必须大于 0 kHz；异常源行：{rows}")
    if np.allclose(selected["测量值"].to_numpy(), 0.0):
        raise ValueError("全部测量值均为 0，无法拟合校准模型。")
    if len(selected) < 6:
        raise ValueError("至少需要 6 个有效样本，当前样本不足。")

    duplicate_mask = selected.duplicated(["测量值", "真值", "频率"], keep=False)
    duplicate_rows = int(duplicate_mask.sum())
    repeated_conditions = int(selected.duplicated(["真值", "频率"], keep=False).sum())

    selected = selected[["源行号", "测量值", "真值", "频率"]].reset_index(drop=True)
    audit = {
        "input_path": str(path),
        "input_sha256": file_sha256(path),
        "sheet": sheet,
        "source_columns": {key: str(value) for key, value in mapping.items()},
        "blank_rows_removed": blank_rows_removed,
        "exact_duplicate_rows": duplicate_rows,
        "repeated_condition_rows": repeated_conditions,
        "sample_count": int(len(selected)),
        "frequency_count": int(selected["频率"].nunique()),
    }
    return selected, audit


# --------------------------- 物理约束模型 ---------------------------
@dataclass(frozen=True)
class Term:
    name: str
    frequency_power: int = 0
    amplitude_power: int = 0

    @property
    def factor_label(self) -> str:
        factors: list[str] = []
        if self.frequency_power == 1:
            factors.append("u")
        elif self.frequency_power > 1:
            factors.append(f"u^{self.frequency_power}")
        if self.amplitude_power == 2:
            factors.append("z^2")
        elif self.amplitude_power > 0:
            factors.append(f"z^{self.amplitude_power}")
        return "·".join(factors) if factors else "1"


@dataclass(frozen=True)
class ModelSpec:
    key: str
    name: str
    terms: tuple[Term, ...]
    description: str
    kind: str = "basis"

    @property
    def parameter_count(self) -> int:
        return len(self.terms)

    @property
    def is_identity(self) -> bool:
        return self.kind == "identity"

    @property
    def is_lookup(self) -> bool:
        return self.kind == "lookup"

    @property
    def is_hybrid_lookup(self) -> bool:
        return self.kind == "hybrid_lookup"


G0 = Term("g0")
GF1 = Term("g_f1", frequency_power=1)
GF2 = Term("g_f2", frequency_power=2)
GA2 = Term("g_a2", amplitude_power=2)
GA2F1 = Term("g_a2_f1", frequency_power=1, amplitude_power=2)
GA2F2 = Term("g_a2_f2", frequency_power=2, amplitude_power=2)
GA4 = Term("g_a4", amplitude_power=4)

MODELS: tuple[ModelSpec, ...] = (
    ModelSpec("identity", "未校准基线", (), "Vcal = Vmeas", "identity"),
    ModelSpec("global_gain", "全局比例", (G0,), "单一比例增益"),
    ModelSpec(
        "frequency_gain_lookup",
        "分频率增益线性插值",
        (),
        "各校准频率独立估计增益，频率结点之间线性插值",
        "lookup",
    ),
    ModelSpec(
        "frequency_gain_amplitude_quintic",
        "分频率增益+平滑幅值修正",
        (),
        "分频率增益线性插值，并以共享三次、五次奇函数修正小幅非线性",
        "hybrid_lookup",
    ),
    ModelSpec(
        "linear_frequency_gain",
        "频率一次增益",
        (G0, GF1),
        "增益随频率一次变化",
    ),
    ModelSpec(
        "quadratic_frequency_gain",
        "频率二次增益",
        (G0, GF1, GF2),
        "增益随频率二次变化",
    ),
    ModelSpec(
        "cubic_amplitude",
        "三次幅值修正",
        (G0, GA2),
        "等效瞬时值奇三次修正",
    ),
    ModelSpec(
        "cubic_frequency_linear",
        "三次幅值+频率一次",
        (G0, GF1, GA2, GA2F1),
        "线性频响与三次非线性的联合模型",
    ),
    ModelSpec(
        "quintic_frequency_linear",
        "五次幅值+频率一次",
        (G0, GF1, GA2, GA2F1, GA4),
        "线性频响、三次和五次非线性的联合模型",
    ),
    ModelSpec(
        "cubic_frequency_quadratic",
        "三次幅值+频率二次",
        (G0, GF1, GF2, GA2, GA2F1, GA2F2),
        "二次频响与三次非线性的联合模型",
    ),
)


@dataclass(frozen=True)
class Normalization:
    voltage_scale_mv: float
    frequency_center_khz: float
    frequency_scale_khz: float


def make_normalization(vpp_meas: np.ndarray, frequency: np.ndarray) -> Normalization:
    voltage_scale = float(np.max(np.abs(vpp_meas)))
    frequency_min = float(np.min(frequency))
    frequency_max = float(np.max(frequency))
    frequency_center = (frequency_min + frequency_max) / 2.0
    frequency_scale = (frequency_max - frequency_min) / 2.0
    if voltage_scale <= EPS:
        raise ValueError("测量值尺度为 0，无法归一化。")
    if frequency_scale <= EPS:
        frequency_scale = 1.0
    return Normalization(voltage_scale, frequency_center, frequency_scale)


def design_matrix(
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    spec: ModelSpec,
    normalization: Normalization,
) -> np.ndarray:
    x = np.asarray(vpp_meas, dtype=float)
    f = np.asarray(frequency, dtype=float)
    if spec.is_identity:
        return np.empty((len(x), 0), dtype=float)
    if spec.is_lookup or spec.is_hybrid_lookup:
        raise ValueError("分频率查表模型没有固定设计矩阵。")
    u = (f - normalization.frequency_center_khz) / normalization.frequency_scale_khz
    z = x / normalization.voltage_scale_mv
    columns = [
        x * np.power(u, term.frequency_power) * np.power(z, term.amplitude_power)
        for term in spec.terms
    ]
    return np.column_stack(columns)


def fit_coefficients(
    spec: ModelSpec,
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    vpp_true: np.ndarray,
    normalization: Normalization,
) -> tuple[np.ndarray, int, float]:
    if spec.is_identity:
        return np.empty(0, dtype=float), 0, 1.0
    if spec.is_lookup:
        knots = np.sort(np.unique(frequency))
        if len(knots) < 2:
            raise ValueError("分频率插值模型至少需要两个不同频率。")
        gains = np.empty(len(knots), dtype=float)
        energy = np.empty(len(knots), dtype=float)
        for index, knot in enumerate(knots):
            selected = frequency == knot
            x_group = np.asarray(vpp_meas[selected], dtype=float)
            y_group = np.asarray(vpp_true[selected], dtype=float)
            denominator = float(x_group @ x_group)
            if denominator <= EPS:
                raise ValueError(f"{knot:g} kHz 组没有非零测量值。")
            gains[index] = float(x_group @ y_group) / denominator
            energy[index] = denominator
        packed = np.empty(2 * len(knots), dtype=float)
        packed[0::2] = knots
        packed[1::2] = gains
        condition = float(np.sqrt(np.max(energy) / np.min(energy)))
        return packed, int(len(knots)), condition
    if spec.is_hybrid_lookup:
        knots = np.sort(np.unique(frequency))
        if len(knots) < 2:
            raise ValueError("混合查表模型至少需要两个不同频率。")
        z = vpp_meas / normalization.voltage_scale_mv
        matrix = np.zeros((len(vpp_meas), len(knots) + 2), dtype=float)
        for index, knot in enumerate(knots):
            selected = frequency == knot
            matrix[selected, index] = vpp_meas[selected]
        matrix[:, -2] = vpp_meas * z**2
        matrix[:, -1] = vpp_meas * z**4
        solved, _, rank, singular_values = np.linalg.lstsq(matrix, vpp_true, rcond=None)
        if rank < matrix.shape[1]:
            raise ValueError("混合查表模型设计矩阵秩不足。")
        packed = np.concatenate(
            ([float(len(knots))], knots, solved[: len(knots)], solved[-2:])
        )
        condition = float(singular_values[0] / singular_values[-1])
        return packed, int(rank), condition
    matrix = design_matrix(vpp_meas, frequency, spec, normalization)
    if len(vpp_true) <= spec.parameter_count:
        raise ValueError("样本数不足以估计模型参数。")
    coefficients, _, rank, singular_values = np.linalg.lstsq(
        matrix, vpp_true, rcond=None
    )
    if rank < spec.parameter_count:
        raise ValueError(
            f"设计矩阵秩不足（{rank}/{spec.parameter_count}），当前数据无法辨识该模型。"
        )
    condition = (
        float(singular_values[0] / singular_values[-1])
        if singular_values[-1] > EPS
        else math.inf
    )
    if not np.isfinite(coefficients).all():
        raise ValueError("拟合得到非有限系数。")
    return coefficients, int(rank), condition


def lookup_gain(
    frequency: np.ndarray,
    packed_coefficients: np.ndarray,
) -> np.ndarray:
    """结点间线性插值；交叉验证时在两端按最近线段外推。"""
    target = np.asarray(frequency, dtype=float)
    knots = np.asarray(packed_coefficients[0::2], dtype=float)
    gains = np.asarray(packed_coefficients[1::2], dtype=float)
    if len(knots) < 2:
        raise ValueError("线性插值至少需要两个频率结点。")
    interpolated = np.interp(target, knots, gains)
    below = target < knots[0]
    above = target > knots[-1]
    if below.any():
        slope = (gains[1] - gains[0]) / (knots[1] - knots[0])
        interpolated[below] = gains[0] + slope * (target[below] - knots[0])
    if above.any():
        slope = (gains[-1] - gains[-2]) / (knots[-1] - knots[-2])
        interpolated[above] = gains[-1] + slope * (target[above] - knots[-1])
    return interpolated


def unpack_hybrid_coefficients(
    packed_coefficients: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, float, float]:
    count = int(round(float(packed_coefficients[0])))
    knots = packed_coefficients[1 : 1 + count]
    gains = packed_coefficients[1 + count : 1 + 2 * count]
    amplitude_cubic, amplitude_quintic = packed_coefficients[-2:]
    return knots, gains, float(amplitude_cubic), float(amplitude_quintic)


def predict(
    spec: ModelSpec,
    coefficients: np.ndarray,
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    normalization: Normalization,
) -> np.ndarray:
    x = np.asarray(vpp_meas, dtype=float)
    if spec.is_identity:
        return x.copy()
    if spec.is_lookup:
        return x * lookup_gain(np.asarray(frequency, dtype=float), coefficients)
    if spec.is_hybrid_lookup:
        knots, gains, cubic, quintic = unpack_hybrid_coefficients(coefficients)
        interleaved = np.empty(2 * len(knots), dtype=float)
        interleaved[0::2] = knots
        interleaved[1::2] = gains
        z = x / normalization.voltage_scale_mv
        return x * (
            lookup_gain(np.asarray(frequency, dtype=float), interleaved)
            + cubic * z**2
            + quintic * z**4
        )
    return design_matrix(x, frequency, spec, normalization) @ coefficients


def coefficient_confidence_intervals(
    matrix: np.ndarray,
    residuals: np.ndarray,
    groups: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, str]:
    """
    计算 95% 系数区间。

    多频率数据采用按频率聚类的稳健协方差；只有一个频率时退化为普通 OLS
    近似区间。区间用于判断系数稳定性，不应被解释为仪器不确定度。
    """
    n, p = matrix.shape
    if p == 0:
        empty = np.empty(0, dtype=float)
        return empty, empty, "无参数"
    bread = np.linalg.pinv(matrix.T @ matrix)
    unique_groups = np.unique(groups)
    if len(unique_groups) >= 2 and n > p:
        meat = np.zeros((p, p), dtype=float)
        for group in unique_groups:
            index = groups == group
            score = matrix[index].T @ residuals[index]
            meat += np.outer(score, score)
        correction = (
            len(unique_groups) / (len(unique_groups) - 1) * (n - 1) / max(n - p, 1)
        )
        covariance = correction * bread @ meat @ bread
        critical = float(student_t.ppf(0.975, df=len(unique_groups) - 1))
        method = f"按频率聚类稳健 95% CI（df={len(unique_groups) - 1}）"
    else:
        degrees = max(n - p, 1)
        sigma2 = float(residuals @ residuals) / degrees
        covariance = sigma2 * bread
        critical = float(student_t.ppf(0.975, df=degrees))
        method = f"OLS 近似 95% CI（df={degrees}）"
    standard_error = np.sqrt(np.clip(np.diag(covariance), 0.0, None))
    return (
        -critical * standard_error,
        critical * standard_error,
        method,
    )


def lookup_coefficient_rows(
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    vpp_true: np.ndarray,
    packed_coefficients: np.ndarray,
) -> tuple[list[dict[str, float | str]], str]:
    """为每个频率结点的过原点增益给出普通 OLS 近似区间。"""
    knots = packed_coefficients[0::2]
    gains = packed_coefficients[1::2]
    rows: list[dict[str, float | str]] = []
    for knot, gain in zip(knots, gains):
        selected = frequency == knot
        x_group = vpp_meas[selected]
        y_group = vpp_true[selected]
        residual = gain * x_group - y_group
        degrees = max(len(x_group) - 1, 1)
        sigma2 = float(residual @ residual) / degrees
        denominator = float(x_group @ x_group)
        standard_error = math.sqrt(max(sigma2 / denominator, 0.0))
        critical = float(student_t.ppf(0.975, df=degrees))
        rows.append(
            {
                "name": f"gain_{knot:g}_kHz",
                "factor": f"f={knot:g} kHz 结点",
                "frequency_khz": float(knot),
                "value": float(gain),
                "ci95_low": float(gain - critical * standard_error),
                "ci95_high": float(gain + critical * standard_error),
            }
        )
    return rows, "各频率过原点 OLS 近似 95% CI"


def hybrid_coefficient_rows(
    packed_coefficients: np.ndarray,
) -> tuple[list[dict[str, float | str]], str]:
    knots, gains, cubic, quintic = unpack_hybrid_coefficients(packed_coefficients)
    rows: list[dict[str, float | str]] = [
        {
            "name": f"gain_{knot:g}_kHz",
            "factor": f"f={knot:g} kHz 基础增益结点",
            "frequency_khz": float(knot),
            "value": float(gain),
            "ci95_low": None,
            "ci95_high": None,
        }
        for knot, gain in zip(knots, gains)
    ]
    rows.extend(
        [
            {
                "name": "g_a2",
                "factor": "z^2（共享）",
                "value": cubic,
                "ci95_low": None,
                "ci95_high": None,
            },
            {
                "name": "g_a4",
                "factor": "z^4（共享）",
                "value": quintic,
                "ci95_low": None,
                "ci95_high": None,
            },
        ]
    )
    return rows, "联合最小二乘点估计；交叉验证用于评估系数整体稳定性"


# --------------------------- 评估与交叉验证 ---------------------------
def calculate_metrics(truth: np.ndarray, prediction: np.ndarray) -> dict[str, Any]:
    truth = np.asarray(truth, dtype=float)
    prediction = np.asarray(prediction, dtype=float)
    error = prediction - truth
    absolute = np.abs(error)
    squared = error**2
    nonzero = np.abs(truth) > EPS
    percent_error = np.full_like(truth, np.nan)
    percent_error[nonzero] = 100.0 * absolute[nonzero] / np.abs(truth[nonzero])
    centered = truth - np.mean(truth)
    total = float(centered @ centered)
    r2 = None if total <= EPS else 1.0 - float(squared.sum()) / total
    return {
        "count": int(len(truth)),
        "bias_mv": float(np.mean(error)),
        "median_error_mv": float(np.median(error)),
        "mae_mv": float(np.mean(absolute)),
        "median_abs_error_mv": float(np.median(absolute)),
        "rmse_mv": float(np.sqrt(np.mean(squared))),
        "p95_abs_error_mv": float(np.percentile(absolute, 95)),
        "p99_abs_error_mv": float(np.percentile(absolute, 99)),
        "max_abs_error_mv": float(np.max(absolute)),
        "r2": r2,
        "mape_percent": (
            float(np.mean(percent_error[nonzero])) if nonzero.any() else None
        ),
        "p95_abs_percent": (
            float(np.percentile(percent_error[nonzero], 95)) if nonzero.any() else None
        ),
        "residual_std_mv": (float(np.std(error, ddof=1)) if len(error) > 1 else 0.0),
        "within_0_5mv_percent": float(100.0 * np.mean(absolute <= 0.5)),
        "within_1mv_percent": float(100.0 * np.mean(absolute <= 1.0)),
        "within_2mv_percent": float(100.0 * np.mean(absolute <= 2.0)),
        "within_5mv_percent": float(100.0 * np.mean(absolute <= 5.0)),
    }


def information_criteria(
    truth: np.ndarray, prediction: np.ndarray, parameter_count: int
) -> dict[str, float | None]:
    n = len(truth)
    error = prediction - truth
    sse = max(float(error @ error), np.finfo(float).tiny)
    log_likelihood_part = n * math.log(sse / n)
    aic = log_likelihood_part + 2 * parameter_count
    denominator = n - parameter_count - 1
    aicc = (
        aic + 2 * parameter_count * (parameter_count + 1) / denominator
        if denominator > 0
        else None
    )
    bic = log_likelihood_part + parameter_count * math.log(n)
    return {"aic": aic, "aicc": aicc, "bic": bic}


def build_cv_splits(
    frequency: np.ndarray,
    mode: str,
    folds: int,
    seed: int,
) -> tuple[list[tuple[np.ndarray, np.ndarray, str]], str]:
    n = len(frequency)
    unique_frequency = np.unique(frequency)
    effective_mode = mode
    if mode == "auto":
        effective_mode = "frequency" if len(unique_frequency) >= 2 else "kfold"
    if effective_mode == "none":
        return [], "未执行交叉验证"
    if effective_mode == "frequency":
        if len(unique_frequency) < 2:
            raise ValueError("按频率交叉验证至少需要两个不同频率。")
        splits = []
        for value in np.sort(unique_frequency):
            test = np.flatnonzero(frequency == value)
            train = np.flatnonzero(frequency != value)
            splits.append((train, test, f"{value:g} kHz"))
        return splits, f"按频率留一组交叉验证（{len(splits)} 折）"
    if effective_mode == "kfold":
        fold_count = min(max(int(folds), 2), n)
        if fold_count < 2:
            raise ValueError("K 折交叉验证至少需要 2 折。")
        indices = np.arange(n)
        rng = np.random.default_rng(seed)
        rng.shuffle(indices)
        test_groups = np.array_split(indices, fold_count)
        splits = []
        for number, test in enumerate(test_groups, start=1):
            train = np.setdiff1d(indices, test, assume_unique=True)
            splits.append((train, np.sort(test), f"第 {number} 折"))
        return splits, f"随机 K 折交叉验证（{fold_count} 折，seed={seed}）"
    raise ValueError(f"未知交叉验证模式：{mode}")


def cross_validate_model(
    spec: ModelSpec,
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    vpp_true: np.ndarray,
    normalization: Normalization,
    splits: Sequence[tuple[np.ndarray, np.ndarray, str]],
) -> tuple[np.ndarray | None, list[dict[str, Any]]]:
    if not splits:
        return None, []
    out_of_fold = np.full(len(vpp_true), np.nan, dtype=float)
    fold_rows: list[dict[str, Any]] = []
    for fold_number, (train, test, label) in enumerate(splits, start=1):
        coefficients, _, _ = fit_coefficients(
            spec,
            vpp_meas[train],
            frequency[train],
            vpp_true[train],
            normalization,
        )
        fold_prediction = predict(
            spec,
            coefficients,
            vpp_meas[test],
            frequency[test],
            normalization,
        )
        out_of_fold[test] = fold_prediction
        fold_metrics = calculate_metrics(vpp_true[test], fold_prediction)
        fold_rows.append(
            {
                "fold": fold_number,
                "held_out": label,
                "train_count": int(len(train)),
                "test_count": int(len(test)),
                **fold_metrics,
            }
        )
    if not np.isfinite(out_of_fold).all():
        raise RuntimeError("交叉验证没有覆盖全部样本。")
    return out_of_fold, fold_rows


def evaluate_models(
    vpp_meas: np.ndarray,
    frequency: np.ndarray,
    vpp_true: np.ndarray,
    normalization: Normalization,
    splits: Sequence[tuple[np.ndarray, np.ndarray, str]],
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    results: dict[str, dict[str, Any]] = {}
    skipped: dict[str, str] = {}
    for spec in MODELS:
        try:
            coefficients, rank, condition = fit_coefficients(
                spec, vpp_meas, frequency, vpp_true, normalization
            )
            fitted = predict(spec, coefficients, vpp_meas, frequency, normalization)
            fit_metrics = calculate_metrics(vpp_true, fitted)
            oof, fold_rows = cross_validate_model(
                spec,
                vpp_meas,
                frequency,
                vpp_true,
                normalization,
                splits,
            )
            cv_metrics = calculate_metrics(vpp_true, oof) if oof is not None else None
            # 频率留组验证回答“未校准频点如何插值”，但无法发现幅值高阶项
            # 对标定点的过拟合。再做一次按参考幅值整组留出，两种外推风险分开报告。
            amplitude_splits = []
            for value in np.sort(np.unique(vpp_true)):
                amplitude_test = np.flatnonzero(vpp_true == value)
                amplitude_train = np.flatnonzero(vpp_true != value)
                amplitude_splits.append(
                    (amplitude_train, amplitude_test, f"{value:g} mV Vpp")
                )
            amplitude_oof, amplitude_fold_rows = cross_validate_model(
                spec,
                vpp_meas,
                frequency,
                vpp_true,
                normalization,
                amplitude_splits,
            )
            amplitude_cv_metrics = calculate_metrics(vpp_true, amplitude_oof)
            fold_rmse = np.array([row["rmse_mv"] for row in fold_rows], dtype=float)
            fold_summary = (
                {
                    "rmse_mean_mv": float(np.mean(fold_rmse)),
                    "rmse_std_mv": (
                        float(np.std(fold_rmse, ddof=1)) if len(fold_rmse) > 1 else 0.0
                    ),
                    "rmse_se_mv": (
                        float(np.std(fold_rmse, ddof=1) / np.sqrt(len(fold_rmse)))
                        if len(fold_rmse) > 1
                        else 0.0
                    ),
                }
                if len(fold_rmse)
                else None
            )
            residuals = fitted - vpp_true
            parameter_count = (
                rank
                if spec.is_lookup or spec.is_hybrid_lookup
                else spec.parameter_count
            )
            if spec.is_lookup:
                coefficient_rows, ci_method = lookup_coefficient_rows(
                    vpp_meas, frequency, vpp_true, coefficients
                )
            elif spec.is_hybrid_lookup:
                coefficient_rows, ci_method = hybrid_coefficient_rows(coefficients)
            else:
                matrix = design_matrix(vpp_meas, frequency, spec, normalization)
                ci_low_delta, ci_high_delta, ci_method = (
                    coefficient_confidence_intervals(matrix, residuals, frequency)
                )
                coefficient_rows = []
                for index, term in enumerate(spec.terms):
                    coefficient_rows.append(
                        {
                            "name": term.name,
                            "factor": term.factor_label,
                            "value": float(coefficients[index]),
                            "ci95_low": float(
                                coefficients[index] + ci_low_delta[index]
                            ),
                            "ci95_high": float(
                                coefficients[index] + ci_high_delta[index]
                            ),
                        }
                    )
            results[spec.key] = {
                "spec": spec,
                "parameter_count": parameter_count,
                "coefficients": coefficients,
                "coefficient_rows": coefficient_rows,
                "coefficient_ci_method": ci_method,
                "rank": rank,
                "condition_number": condition,
                "fitted_prediction": fitted,
                "fit_metrics": fit_metrics,
                "oof_prediction": oof,
                "cv_metrics": cv_metrics,
                "amplitude_oof_prediction": amplitude_oof,
                "amplitude_cv_metrics": amplitude_cv_metrics,
                "amplitude_fold_rows": amplitude_fold_rows,
                "fold_rows": fold_rows,
                "fold_summary": fold_summary,
                **information_criteria(vpp_true, fitted, parameter_count),
            }
        except Exception as exc:
            skipped[spec.name] = str(exc)
    return results, skipped


def select_recommended_model(
    results: dict[str, dict[str, Any]],
) -> tuple[str, dict[str, Any]]:
    if not results:
        raise RuntimeError("所有候选模型均拟合失败。")
    with_cv = [
        (key, value)
        for key, value in results.items()
        if value["fold_summary"] is not None
    ]
    if with_cv:
        # 主排序仍采用更苛刻的频率留组误差；幅值留组 RMSE 作为第二目标，
        # 防止高阶非线性项仅仅记住本表的 12 个幅值等级。
        best_key, best = min(
            with_cv,
            key=lambda item: (
                item[1]["fold_summary"]["rmse_mean_mv"],
                item[1]["amplitude_cv_metrics"]["rmse_mv"],
            ),
        )
        threshold = (
            best["fold_summary"]["rmse_mean_mv"] + best["fold_summary"]["rmse_se_mv"]
        )
        eligible = [
            (key, value)
            for key, value in with_cv
            if value["fold_summary"]["rmse_mean_mv"] <= threshold + 1e-12
        ]
        recommended_key, recommended = min(
            eligible,
            key=lambda item: (
                item[1]["parameter_count"],
                item[1]["fold_summary"]["rmse_mean_mv"],
            ),
        )
        baseline_lookup = results.get("frequency_gain_lookup")
        hybrid_lookup = results.get("frequency_gain_amplitude_quintic")
        hybrid_selected = False
        if baseline_lookup is not None and hybrid_lookup is not None:
            fit_ratio = (
                hybrid_lookup["fit_metrics"]["rmse_mv"]
                / baseline_lookup["fit_metrics"]["rmse_mv"]
            )
            frequency_ratio = (
                hybrid_lookup["cv_metrics"]["rmse_mv"]
                / baseline_lookup["cv_metrics"]["rmse_mv"]
            )
            amplitude_ratio = (
                hybrid_lookup["amplitude_cv_metrics"]["rmse_mv"]
                / baseline_lookup["amplitude_cv_metrics"]["rmse_mv"]
            )
            # 只有训练误差显著下降，且两种留组验证均没有实质退化时，
            # 才允许两个共享幅值项替代纯增益查表。
            if (
                fit_ratio <= 0.80
                and frequency_ratio <= 1.02
                and amplitude_ratio <= 1.02
            ):
                recommended_key = "frequency_gain_amplitude_quintic"
                recommended = hybrid_lookup
                hybrid_selected = True
        selection = {
            "criterion": "频率留组 CV 一标准误差规则，并以幅值留组 CV 防止过拟合",
            "lowest_cv_model": best["spec"].name,
            "lowest_cv_fold_rmse_mean_mv": best["fold_summary"]["rmse_mean_mv"],
            "one_standard_error_threshold_mv": threshold,
            "eligible_models": [value["spec"].name for _, value in eligible],
            "amplitude_holdout_rmse_mv": recommended["amplitude_cv_metrics"]["rmse_mv"],
            "hybrid_pareto_override": hybrid_selected,
            "reason": (
                "先找平均验证 RMSE 最低的模型，再在其一个标准误差范围内"
                "选择参数最少者；同时整组留出每个参考幅值，检查幅值非线性项"
                "是否只是在记忆校准表。若共享幅值修正使拟合 RMSE 至少降低 20%，"
                "且两类留组 RMSE 均不恶化超过 2%，则采用该修正。"
            ),
        }
        return recommended_key, selection

    recommended_key, recommended = min(
        results.items(),
        key=lambda item: (
            item[1]["bic"],
            item[1]["parameter_count"],
        ),
    )
    return recommended_key, {
        "criterion": "BIC（未执行交叉验证）",
        "eligible_models": [recommended["spec"].name],
        "reason": "未执行交叉验证，按 BIC 在拟合误差与模型复杂度之间折中。",
    }


# --------------------------- 公式、单调性和部署导出 ---------------------------
def signed_expression(
    coefficients: np.ndarray, terms: Sequence[Term], variable: str = "z"
) -> str:
    if len(coefficients) == 0:
        return "1"
    parts: list[str] = []
    for index, (coefficient, term) in enumerate(zip(coefficients, terms)):
        factor = term.factor_label.replace("z", variable).replace("^", "**")
        magnitude = abs(float(coefficient))
        body = f"{magnitude:.12g}"
        if factor != "1":
            body += f"*{factor}"
        if index == 0:
            parts.append(("-" if coefficient < 0 else "") + body)
        else:
            parts.append((" - " if coefficient < 0 else " + ") + body)
    return "".join(parts)


def model_formula(
    spec: ModelSpec,
    coefficients: np.ndarray,
    normalization: Normalization,
) -> dict[str, str]:
    if spec.is_lookup:
        knots = coefficients[0::2]
        gains = coefficients[1::2]
        pairs = ", ".join(
            f"{frequency:g}:{gain:.12g}" for frequency, gain in zip(knots, gains)
        )
        return {
            "definitions": ("g(f) 在以下 f_kHz:gain 结点间作分段线性插值：" + pairs),
            "vpp": "Vcal = Vmeas * g(f_kHz)",
            "instantaneous": "vout = vin * g(f_kHz)",
            "note": (
                "该模型只校正幅值增益；当前 Vpp 表格不能验证相位、直流偏置"
                "或逐采样点波形失真。"
            ),
        }
    if spec.is_hybrid_lookup:
        knots, gains, cubic, quintic = unpack_hybrid_coefficients(coefficients)
        pairs = ", ".join(
            f"{frequency:g}:{gain:.12g}" for frequency, gain in zip(knots, gains)
        )
        return {
            "definitions": (
                "g(f) 在频率结点间线性插值；"
                f"z=Vmeas/{normalization.voltage_scale_mv:.12g}；"
                f"g_a2={cubic:.12g}，g_a4={quintic:.12g}；结点：{pairs}"
            ),
            "vpp": "Vcal = Vmeas * (g(f_kHz) + g_a2*z^2 + g_a4*z^4)",
            "instantaneous": "vout = vin * (g(f_kHz) + g_a2*q^2 + g_a4*q^4)，q=2*vin/Vscale",
            "note": "共享幅值项保持零点对称；当前 Vpp 数据不能验证相位或波形失真。",
        }
    bracket = signed_expression(coefficients, spec.terms)
    if spec.is_identity:
        bracket = "1"
    definitions = (
        f"u=(f_kHz-{normalization.frequency_center_khz:.12g})/"
        f"{normalization.frequency_scale_khz:.12g}, "
        f"z=Vmeas/{normalization.voltage_scale_mv:.12g}"
    )
    return {
        "definitions": definitions,
        "vpp": f"Vcal = Vmeas * ({bracket})",
        "instantaneous": (
            "vout = vin * (同一括号，其中 z 替换为 "
            f"2*vin/{normalization.voltage_scale_mv:.12g})"
        ),
        "note": (
            "瞬时值等效式仅保证在正弦、零偏置且校准函数于有效范围单调时"
            "复现 Vpp 校准；当前表格数据本身不能验证完整波形失真。"
        ),
    }


def check_instantaneous_monotonicity(
    spec: ModelSpec,
    coefficients: np.ndarray,
    normalization: Normalization,
    voltage_max_mv: float,
    frequency_min_khz: float,
    frequency_max_khz: float,
) -> dict[str, Any]:
    if spec.is_identity:
        return {
            "monotonic_on_calibration_domain": True,
            "minimum_derivative": 1.0,
        }
    if spec.is_lookup:
        gains = coefficients[1::2]
        minimum = float(np.min(gains))
        return {
            "monotonic_on_calibration_domain": bool(minimum > 0.0),
            "minimum_derivative": minimum,
        }
    if spec.is_hybrid_lookup:
        _, gains, cubic, quintic = unpack_hybrid_coefficients(coefficients)
        q = np.linspace(-1.0, 1.0, 1001)
        derivative = (
            gains[:, None]
            + 3.0 * cubic * q[None, :] ** 2
            + 5.0 * quintic * q[None, :] ** 4
        )
        minimum = float(np.min(derivative))
        return {
            "monotonic_on_calibration_domain": bool(minimum > 0.0),
            "minimum_derivative": minimum,
        }
    voltage = np.linspace(-voltage_max_mv / 2.0, voltage_max_mv / 2.0, 401)
    frequency = np.linspace(frequency_min_khz, frequency_max_khz, 201)
    q = 2.0 * voltage[None, :] / normalization.voltage_scale_mv
    u = (
        frequency[:, None] - normalization.frequency_center_khz
    ) / normalization.frequency_scale_khz
    derivative = np.zeros((len(frequency), len(voltage)), dtype=float)
    for coefficient, term in zip(coefficients, spec.terms):
        derivative += (
            coefficient
            * np.power(u, term.frequency_power)
            * (term.amplitude_power + 1)
            * np.power(q, term.amplitude_power)
        )
    minimum = float(np.min(derivative))
    return {
        "monotonic_on_calibration_domain": bool(minimum > 0.0),
        "minimum_derivative": minimum,
    }


def c_factor_expression(
    spec: ModelSpec, coefficients: np.ndarray, amplitude_variable: str
) -> str:
    if spec.is_identity:
        return "1.0f"
    pieces: list[str] = []
    for coefficient, term in zip(coefficients, spec.terms):
        factors = [f"{float(coefficient):.9e}f"]
        if term.frequency_power == 1:
            factors.append("u")
        elif term.frequency_power == 2:
            factors.append("u2")
        if term.amplitude_power == 2:
            factors.append(amplitude_variable)
        elif term.amplitude_power == 4:
            factors.append(f"{amplitude_variable}2")
        body = " * ".join(factors)
        pieces.append(f"({body})")
    return " + ".join(pieces)


def build_lookup_c_header(
    spec: ModelSpec,
    coefficients: np.ndarray,
    ranges: dict[str, float],
    source_hash: str,
) -> str:
    knots = coefficients[0::2]
    gains = coefficients[1::2]
    knot_values = ", ".join(f"{value:.9e}f" for value in knots)
    gain_values = ", ".join(f"{value:.9e}f" for value in gains)
    return f"""/*
 * 自动生成的 Vpp 校准模型
 * 模型：{spec.name}
 * 源数据 SHA-256：{source_hash}
 * 有效范围：Vpp {ranges["voltage_min_mv"]:.9g}~{ranges["voltage_max_mv"]:.9g} mV，
 *           f {ranges["frequency_min_khz"]:.9g}~{ranges["frequency_max_khz"]:.9g} kHz
 *
 * 结点间使用线性插值。范围外未经验证；调用者应先执行范围检查。
 */
#ifndef CB_CALIBRATION_MODEL_H
#define CB_CALIBRATION_MODEL_H

#include <stdbool.h>

#define CB_VOLTAGE_MIN_MV ({ranges["voltage_min_mv"]:.9e}f)
#define CB_VOLTAGE_MAX_MV ({ranges["voltage_max_mv"]:.9e}f)
#define CB_FREQUENCY_MIN_KHZ ({ranges["frequency_min_khz"]:.9e}f)
#define CB_FREQUENCY_MAX_KHZ ({ranges["frequency_max_khz"]:.9e}f)
#define CB_GAIN_KNOT_COUNT ({len(knots)}u)

static const float CB_FREQUENCY_KNOTS_KHZ[CB_GAIN_KNOT_COUNT] = {{
    {knot_values}
}};
static const float CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT] = {{
    {gain_values}
}};

static inline bool cb_calibration_in_range(float vpp_mv, float frequency_khz)
{{
    return vpp_mv >= CB_VOLTAGE_MIN_MV &&
           vpp_mv <= CB_VOLTAGE_MAX_MV &&
           frequency_khz >= CB_FREQUENCY_MIN_KHZ &&
           frequency_khz <= CB_FREQUENCY_MAX_KHZ;
}}

static inline float cb_interpolated_gain(float frequency_khz)
{{
    unsigned int i;
    if (frequency_khz <= CB_FREQUENCY_KNOTS_KHZ[0]) {{
        return CB_GAIN_KNOTS[0];
    }}
    if (frequency_khz >= CB_FREQUENCY_KNOTS_KHZ[CB_GAIN_KNOT_COUNT - 1u]) {{
        return CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT - 1u];
    }}
    for (i = 1u; i < CB_GAIN_KNOT_COUNT; ++i) {{
        if (frequency_khz <= CB_FREQUENCY_KNOTS_KHZ[i]) {{
            const float span = CB_FREQUENCY_KNOTS_KHZ[i] -
                               CB_FREQUENCY_KNOTS_KHZ[i - 1u];
            const float alpha = (frequency_khz -
                                 CB_FREQUENCY_KNOTS_KHZ[i - 1u]) / span;
            return CB_GAIN_KNOTS[i - 1u] +
                   alpha * (CB_GAIN_KNOTS[i] - CB_GAIN_KNOTS[i - 1u]);
        }}
    }}
    return CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT - 1u];
}}

static inline float cb_calibrate_vpp_mv(float vpp_mv, float frequency_khz)
{{
    return vpp_mv * cb_interpolated_gain(frequency_khz);
}}

/*
 * 该等效瞬时值形式只校正幅度增益。
 * 校准表没有验证相位、直流偏置或逐采样点波形失真。
 */
static inline float cb_calibrate_sample_mv(float vin_mv, float frequency_khz)
{{
    return vin_mv * cb_interpolated_gain(frequency_khz);
}}

#endif /* CB_CALIBRATION_MODEL_H */
"""


def build_c_header(
    spec: ModelSpec,
    coefficients: np.ndarray,
    normalization: Normalization,
    ranges: dict[str, float],
    source_hash: str,
) -> str:
    if spec.is_lookup:
        return build_lookup_c_header(spec, coefficients, ranges, source_hash)
    if spec.is_hybrid_lookup:
        knots, gains, cubic, quintic = unpack_hybrid_coefficients(coefficients)
        knot_values = ", ".join(f"{value:.9e}f" for value in knots)
        gain_values = ", ".join(f"{value:.9e}f" for value in gains)
        return f"""/*
 * 自动生成的 Vpp 校准模型
 * 模型：{spec.name}
 * 源数据 SHA-256：{source_hash}
 * 有效范围：Vpp {ranges["voltage_min_mv"]:.9g}~{ranges["voltage_max_mv"]:.9g} mV，
 *           f {ranges["frequency_min_khz"]:.9g}~{ranges["frequency_max_khz"]:.9g} kHz
 * 接口与旧版保持一致；signal.c 可直接替换使用。
 */
#ifndef CB_CALIBRATION_MODEL_H
#define CB_CALIBRATION_MODEL_H

#include <stdbool.h>

#define CB_VOLTAGE_MIN_MV ({ranges["voltage_min_mv"]:.9e}f)
#define CB_VOLTAGE_MAX_MV ({ranges["voltage_max_mv"]:.9e}f)
#define CB_FREQUENCY_MIN_KHZ ({ranges["frequency_min_khz"]:.9e}f)
#define CB_FREQUENCY_MAX_KHZ ({ranges["frequency_max_khz"]:.9e}f)
#define CB_VOLTAGE_SCALE_MV ({normalization.voltage_scale_mv:.9e}f)
#define CB_GAIN_KNOT_COUNT ({len(knots)}u)
#define CB_AMPLITUDE_CUBIC ({cubic:.9e}f)
#define CB_AMPLITUDE_QUINTIC ({quintic:.9e}f)

static const float CB_FREQUENCY_KNOTS_KHZ[CB_GAIN_KNOT_COUNT] = {{
    {knot_values}
}};
static const float CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT] = {{
    {gain_values}
}};

static inline bool cb_calibration_in_range(float vpp_mv, float frequency_khz)
{{
    return vpp_mv >= CB_VOLTAGE_MIN_MV &&
           vpp_mv <= CB_VOLTAGE_MAX_MV &&
           frequency_khz >= CB_FREQUENCY_MIN_KHZ &&
           frequency_khz <= CB_FREQUENCY_MAX_KHZ;
}}

static inline float cb_interpolated_gain(float frequency_khz)
{{
    unsigned int i;
    if (frequency_khz <= CB_FREQUENCY_KNOTS_KHZ[0]) {{
        return CB_GAIN_KNOTS[0];
    }}
    if (frequency_khz >= CB_FREQUENCY_KNOTS_KHZ[CB_GAIN_KNOT_COUNT - 1u]) {{
        return CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT - 1u];
    }}
    for (i = 1u; i < CB_GAIN_KNOT_COUNT; ++i) {{
        if (frequency_khz <= CB_FREQUENCY_KNOTS_KHZ[i]) {{
            const float alpha = (frequency_khz - CB_FREQUENCY_KNOTS_KHZ[i - 1u]) /
                                (CB_FREQUENCY_KNOTS_KHZ[i] - CB_FREQUENCY_KNOTS_KHZ[i - 1u]);
            return CB_GAIN_KNOTS[i - 1u] +
                   alpha * (CB_GAIN_KNOTS[i] - CB_GAIN_KNOTS[i - 1u]);
        }}
    }}
    return CB_GAIN_KNOTS[CB_GAIN_KNOT_COUNT - 1u];
}}

static inline float cb_calibrate_vpp_mv(float vpp_mv, float frequency_khz)
{{
    const float z = vpp_mv / CB_VOLTAGE_SCALE_MV;
    const float z2 = z * z;
    const float gain = cb_interpolated_gain(frequency_khz) +
                       CB_AMPLITUDE_CUBIC * z2 +
                       CB_AMPLITUDE_QUINTIC * z2 * z2;
    return vpp_mv * gain;
}}

static inline float cb_calibrate_sample_mv(float vin_mv, float frequency_khz)
{{
    const float q = (2.0f * vin_mv) / CB_VOLTAGE_SCALE_MV;
    const float q2 = q * q;
    const float gain = cb_interpolated_gain(frequency_khz) +
                       CB_AMPLITUDE_CUBIC * q2 +
                       CB_AMPLITUDE_QUINTIC * q2 * q2;
    return vin_mv * gain;
}}

#endif /* CB_CALIBRATION_MODEL_H */
"""
    vpp_expression = c_factor_expression(spec, coefficients, "z2")
    sample_expression = c_factor_expression(spec, coefficients, "q2")
    return f"""/*
 * 自动生成的 Vpp 校准模型
 * 模型：{spec.name}
 * 源数据 SHA-256：{source_hash}
 * 有效范围：Vpp {ranges["voltage_min_mv"]:.9g}~{ranges["voltage_max_mv"]:.9g} mV，
 *           f {ranges["frequency_min_khz"]:.9g}~{ranges["frequency_max_khz"]:.9g} kHz
 *
 * 注意：范围外属于未经验证的外推，本文件不自动钳位输入。
 */
#ifndef CB_CALIBRATION_MODEL_H
#define CB_CALIBRATION_MODEL_H

#include <stdbool.h>

#define CB_VOLTAGE_MIN_MV ({ranges["voltage_min_mv"]:.9e}f)
#define CB_VOLTAGE_MAX_MV ({ranges["voltage_max_mv"]:.9e}f)
#define CB_FREQUENCY_MIN_KHZ ({ranges["frequency_min_khz"]:.9e}f)
#define CB_FREQUENCY_MAX_KHZ ({ranges["frequency_max_khz"]:.9e}f)
#define CB_VOLTAGE_SCALE_MV ({normalization.voltage_scale_mv:.9e}f)
#define CB_FREQUENCY_CENTER_KHZ ({normalization.frequency_center_khz:.9e}f)
#define CB_FREQUENCY_SCALE_KHZ ({normalization.frequency_scale_khz:.9e}f)

static inline bool cb_calibration_in_range(float vpp_mv, float frequency_khz)
{{
    return vpp_mv >= CB_VOLTAGE_MIN_MV &&
           vpp_mv <= CB_VOLTAGE_MAX_MV &&
           frequency_khz >= CB_FREQUENCY_MIN_KHZ &&
           frequency_khz <= CB_FREQUENCY_MAX_KHZ;
}}

static inline float cb_calibrate_vpp_mv(float vpp_mv, float frequency_khz)
{{
    const float u = (frequency_khz - CB_FREQUENCY_CENTER_KHZ) /
                    CB_FREQUENCY_SCALE_KHZ;
    const float u2 = u * u;
    const float z = vpp_mv / CB_VOLTAGE_SCALE_MV;
    const float z2 = z * z;
    const float z22 = z2 * z2;
    (void)u2;
    (void)z22;
    return vpp_mv * ({vpp_expression});
}}

/*
 * 正弦/零偏置/单调假设下的等效瞬时值形式。
 * 校准表只验证了 Vpp，没有验证逐采样点的波形保真度。
 */
static inline float cb_calibrate_sample_mv(float vin_mv, float frequency_khz)
{{
    const float u = (frequency_khz - CB_FREQUENCY_CENTER_KHZ) /
                    CB_FREQUENCY_SCALE_KHZ;
    const float u2 = u * u;
    const float q = (2.0f * vin_mv) / CB_VOLTAGE_SCALE_MV;
    const float q2 = q * q;
    const float q22 = q2 * q2;
    (void)u2;
    (void)q22;
    return vin_mv * ({sample_expression});
}}

#endif /* CB_CALIBRATION_MODEL_H */
"""


# --------------------------- 结果整理 ---------------------------
def model_comparison_dataframe(
    results: dict[str, dict[str, Any]],
    recommended_key: str,
    raw_rmse: float,
) -> pd.DataFrame:
    rows = []
    for key, result in results.items():
        spec = result["spec"]
        fit = result["fit_metrics"]
        cv = result["cv_metrics"] or {}
        amplitude_cv = result["amplitude_cv_metrics"]
        fold = result["fold_summary"] or {}
        reference_rmse = cv.get("rmse_mv", fit["rmse_mv"])
        rows.append(
            {
                "推荐": "是" if key == recommended_key else "",
                "模型": spec.name,
                "模型标识": spec.key,
                "参数数量": result["parameter_count"],
                "拟合偏差_mV": fit["bias_mv"],
                "拟合MAE_mV": fit["mae_mv"],
                "拟合RMSE_mV": fit["rmse_mv"],
                "拟合P95绝对误差_mV": fit["p95_abs_error_mv"],
                "拟合最大绝对误差_mV": fit["max_abs_error_mv"],
                "拟合R2": fit["r2"],
                "CV偏差_mV": cv.get("bias_mv"),
                "CV_MAE_mV": cv.get("mae_mv"),
                "CV_RMSE_mV": cv.get("rmse_mv"),
                "CV_P95绝对误差_mV": cv.get("p95_abs_error_mv"),
                "CV_P99绝对误差_mV": cv.get("p99_abs_error_mv"),
                "CV最大绝对误差_mV": cv.get("max_abs_error_mv"),
                "CV_MAPE_percent": cv.get("mape_percent"),
                "CV误差小于等于0.5mV_percent": cv.get("within_0_5mv_percent"),
                "CV误差小于等于1mV_percent": cv.get("within_1mv_percent"),
                "CV误差小于等于2mV_percent": cv.get("within_2mv_percent"),
                "CV误差小于等于5mV_percent": cv.get("within_5mv_percent"),
                "CV_R2": cv.get("r2"),
                "幅值留组CV_RMSE_mV": amplitude_cv["rmse_mv"],
                "幅值留组CV_P95绝对误差_mV": amplitude_cv["p95_abs_error_mv"],
                "折RMSE均值_mV": fold.get("rmse_mean_mv"),
                "折RMSE标准误_mV": fold.get("rmse_se_mv"),
                "相对未校准RMSE改善_percent": safe_percent(
                    raw_rmse - reference_rmse, raw_rmse
                ),
                "AICc": result["aicc"],
                "BIC": result["bic"],
                "设计矩阵条件数": result["condition_number"],
            }
        )
    frame = pd.DataFrame(rows)
    sort_column = "折RMSE均值_mV" if frame["折RMSE均值_mV"].notna().any() else "BIC"
    return frame.sort_values([sort_column, "参数数量"], na_position="last").reset_index(
        drop=True
    )


def frequency_summary_dataframe(
    data: pd.DataFrame,
    fitted_prediction: np.ndarray,
    validation_prediction: np.ndarray,
) -> pd.DataFrame:
    rows = []
    truth = data["真值"].to_numpy(dtype=float)
    measured = data["测量值"].to_numpy(dtype=float)
    frequency = data["频率"].to_numpy(dtype=float)
    for value in np.sort(np.unique(frequency)):
        index = frequency == value
        raw = calculate_metrics(truth[index], measured[index])
        fitted = calculate_metrics(truth[index], fitted_prediction[index])
        validation = calculate_metrics(truth[index], validation_prediction[index])
        rows.append(
            {
                "频率_kHz": value,
                "样本数": int(index.sum()),
                "真值最小_mV": float(np.min(truth[index])),
                "真值最大_mV": float(np.max(truth[index])),
                "原始偏差_mV": raw["bias_mv"],
                "原始MAE_mV": raw["mae_mv"],
                "原始RMSE_mV": raw["rmse_mv"],
                "拟合偏差_mV": fitted["bias_mv"],
                "拟合MAE_mV": fitted["mae_mv"],
                "拟合RMSE_mV": fitted["rmse_mv"],
                "CV偏差_mV": validation["bias_mv"],
                "CV_MAE_mV": validation["mae_mv"],
                "CV_RMSE_mV": validation["rmse_mv"],
                "CV_P95绝对误差_mV": validation["p95_abs_error_mv"],
                "CV最大绝对误差_mV": validation["max_abs_error_mv"],
            }
        )
    return pd.DataFrame(rows)


def detect_outliers(
    residuals: np.ndarray, groups: np.ndarray
) -> tuple[np.ndarray, dict[str, Any]]:
    """在每个频率组内部检查验证残差，避免把频率系统偏差误报成坏点。"""
    residuals = np.asarray(residuals, dtype=float)
    groups = np.asarray(groups)
    flags = np.zeros(len(residuals), dtype=bool)
    group_rows: list[dict[str, float | int]] = []
    for group in np.sort(np.unique(groups)):
        selected = groups == group
        values = residuals[selected]
        center = float(np.median(values))
        mad = float(np.median(np.abs(values - center)))
        robust_sigma = 1.4826 * mad
        if robust_sigma <= EPS:
            robust_sigma = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
        cutoff = max(3.0 * robust_sigma, 0.5)
        flags[selected] = np.abs(values - center) > cutoff
        group_rows.append(
            {
                "frequency_khz": float(group),
                "median_residual_mv": center,
                "robust_sigma_mv": robust_sigma,
                "cutoff_mv": cutoff,
                "flagged_count": int(flags[selected].sum()),
            }
        )
    return flags, {
        "method": "按频率组：偏离组内中位数 > max(3×MAD标准差, 0.5 mV)",
        "frequency_groups": group_rows,
    }


def diagnose_local_consistency(
    data: pd.DataFrame,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, dict[str, Any]]:
    """
    检查同一频率内部的幅值序列一致性。

    每个频率用 ratio=真值/测量值 的中位数形成稳健参考增益。只有偏离超过
    6 个组内稳健标准差且绝对偏离超过 0.5 mV 才标记，避免把正常量化波动
    当作坏点。这里只提示复核，从不自动删除。
    """
    measured = data["测量值"].to_numpy(dtype=float)
    truth = data["真值"].to_numpy(dtype=float)
    frequency = data["频率"].to_numpy(dtype=float)
    local_residual = np.zeros(len(data), dtype=float)
    local_score = np.zeros(len(data), dtype=float)
    local_flags = np.zeros(len(data), dtype=bool)
    group_rows: list[dict[str, Any]] = []

    for value in np.sort(np.unique(frequency)):
        selected = frequency == value
        x_group = measured[selected]
        y_group = truth[selected]
        usable = np.abs(x_group) > EPS
        if not usable.any():
            continue
        robust_gain = float(np.median(y_group[usable] / x_group[usable]))
        residual = robust_gain * x_group - y_group
        center = float(np.median(residual))
        mad = float(np.median(np.abs(residual - center)))
        sigma = 1.4826 * mad
        if sigma <= EPS:
            sigma = float(np.std(residual, ddof=1)) if len(residual) > 1 else 0.0
        deviation = np.abs(residual - center)
        score = deviation / sigma if sigma > EPS else np.zeros_like(deviation)
        cutoff_mv = max(6.0 * sigma, 0.5)
        flags = deviation > cutoff_mv
        indices = np.flatnonzero(selected)
        local_residual[indices] = residual
        local_score[indices] = score
        local_flags[indices] = flags
        group_rows.append(
            {
                "frequency_khz": float(value),
                "robust_gain": robust_gain,
                "robust_sigma_mv": sigma,
                "cutoff_mv": cutoff_mv,
                "flagged_count": int(flags.sum()),
            }
        )

    suspects: list[dict[str, Any]] = []
    for index in np.flatnonzero(local_flags):
        suspects.append(
            {
                "source_row": int(data.iloc[index]["源行号"]),
                "frequency_khz": float(frequency[index]),
                "measured_mv": float(measured[index]),
                "reference_mv": float(truth[index]),
                "local_residual_mv": float(local_residual[index]),
                "robust_score": float(local_score[index]),
            }
        )
    return (
        local_residual,
        local_score,
        local_flags,
        {
            "method": ("同频率稳健中位增益；偏离 > max(6×组内MAD标准差, 0.5 mV)"),
            "flagged_count": int(local_flags.sum()),
            "suspects": suspects,
            "frequency_groups": group_rows,
        },
    )


def build_detail_dataframes(
    data: pd.DataFrame,
    results: dict[str, dict[str, Any]],
    recommended_key: str,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    recommended = results[recommended_key]
    truth = data["真值"].to_numpy(dtype=float)
    measured = data["测量值"].to_numpy(dtype=float)
    fitted = recommended["fitted_prediction"]
    validation = (
        recommended["oof_prediction"]
        if recommended["oof_prediction"] is not None
        else fitted
    )
    validation_label = "CV" if recommended["oof_prediction"] is not None else "拟合"
    residuals = validation - truth
    model_flags, model_outlier_info = detect_outliers(
        residuals, data["频率"].to_numpy(dtype=float)
    )
    local_residual, local_score, local_flags, local_info = diagnose_local_consistency(
        data
    )
    nonzero = np.abs(truth) > EPS

    detail = data.copy()
    detail["原始误差_mV"] = measured - truth
    detail["原始绝对百分误差"] = np.where(
        nonzero, 100.0 * np.abs(measured - truth) / np.abs(truth), np.nan
    )
    detail["推荐模型"] = recommended["spec"].name
    detail["整表拟合预测_mV"] = fitted
    detail["整表拟合误差_mV"] = fitted - truth
    detail[f"{validation_label}预测_mV"] = validation
    detail[f"{validation_label}误差_mV"] = residuals
    detail[f"{validation_label}绝对百分误差"] = np.where(
        nonzero, 100.0 * np.abs(residuals) / np.abs(truth), np.nan
    )
    detail["CV残差稳健3sigma疑似异常点"] = np.where(model_flags, "是", "")
    detail["同频局部一致性残差_mV"] = local_residual
    detail["同频局部稳健异常分数"] = local_score
    detail["同频局部一致性疑似异常点"] = np.where(local_flags, "是", "")

    all_predictions = data.copy()
    for result in results.values():
        name = result["spec"].name
        all_predictions[f"{name}_拟合预测_mV"] = result["fitted_prediction"]
        all_predictions[f"{name}_拟合误差_mV"] = result["fitted_prediction"] - truth
        if result["oof_prediction"] is not None:
            all_predictions[f"{name}_CV预测_mV"] = result["oof_prediction"]
            all_predictions[f"{name}_CV误差_mV"] = result["oof_prediction"] - truth
    outlier_info = {
        "cross_validation_residual": {
            **model_outlier_info,
            "flagged_count": int(model_flags.sum()),
        },
        "local_consistency": local_info,
        "flagged_count": int(np.count_nonzero(model_flags | local_flags)),
    }
    return detail, all_predictions, outlier_info


def coefficient_dataframe(
    result: dict[str, Any],
) -> pd.DataFrame:
    rows = []
    for row in result["coefficient_rows"]:
        rows.append(
            {
                "参数名": row["name"],
                "归一化因子": row["factor"],
                "系数": row["value"],
                "95%CI下限": row["ci95_low"],
                "95%CI上限": row["ci95_high"],
                "区间方法": result["coefficient_ci_method"],
            }
        )
    if not rows:
        rows.append(
            {
                "参数名": "无",
                "归一化因子": "1",
                "系数": 1.0,
                "95%CI下限": np.nan,
                "95%CI上限": np.nan,
                "区间方法": "恒等映射",
            }
        )
    return pd.DataFrame(rows)


# --------------------------- 可视化 ---------------------------
def save_figure(
    figure: plt.Figure,
    output_dir: Path,
    stem: str,
    dpi: int,
    legacy_stem: str | None = None,
) -> list[Path]:
    paths = [output_dir / f"{stem}.png", output_dir / f"{stem}.svg"]
    figure.savefig(paths[0], dpi=dpi, bbox_inches="tight")
    figure.savefig(paths[1], bbox_inches="tight")
    if legacy_stem:
        legacy_path = output_dir / f"{legacy_stem}.png"
        figure.savefig(legacy_path, dpi=dpi, bbox_inches="tight")
        paths.append(legacy_path)
    plt.close(figure)
    return paths


def plot_data_overview(data: pd.DataFrame, output_dir: Path, dpi: int) -> list[Path]:
    measured = data["测量值"].to_numpy(dtype=float)
    truth = data["真值"].to_numpy(dtype=float)
    frequency = data["频率"].to_numpy(dtype=float)
    residual = measured - truth
    percent = np.where(np.abs(truth) > EPS, 100.0 * residual / np.abs(truth), np.nan)
    _, _, review_flags, _ = diagnose_local_consistency(data)

    figure, axes = plt.subplots(2, 2, figsize=(13, 9))
    scatter = axes[0, 0].scatter(
        truth, measured, c=frequency, cmap="viridis", alpha=0.78, s=28
    )
    low = min(float(np.min(truth)), float(np.min(measured)))
    high = max(float(np.max(truth)), float(np.max(measured)))
    axes[0, 0].plot([low, high], [low, high], "r--", linewidth=1.2)
    if review_flags.any():
        axes[0, 0].scatter(
            truth[review_flags],
            measured[review_flags],
            marker="x",
            color="red",
            s=90,
            linewidth=2,
            label=plot_text("建议复核", "Review suggested"),
        )
        axes[0, 0].legend()
    axes[0, 0].set(
        xlabel=plot_text("真值 (mV Vpp)", "Reference (mV Vpp)"),
        ylabel=plot_text("测量值 (mV Vpp)", "Measured (mV Vpp)"),
        title=plot_text("测量一致性", "Measured vs reference"),
    )
    figure.colorbar(
        scatter,
        ax=axes[0, 0],
        label=plot_text("频率 (kHz)", "Frequency (kHz)"),
    )

    scatter_residual = axes[0, 1].scatter(
        truth, residual, c=frequency, cmap="viridis", alpha=0.78, s=28
    )
    axes[0, 1].axhline(0, color="r", linestyle="--", linewidth=1.2)
    if review_flags.any():
        axes[0, 1].scatter(
            truth[review_flags],
            residual[review_flags],
            marker="x",
            color="red",
            s=90,
            linewidth=2,
        )
    axes[0, 1].set(
        xlabel=plot_text("真值 (mV Vpp)", "Reference (mV Vpp)"),
        ylabel=plot_text("原始误差 (mV)", "Raw error (mV)"),
        title=plot_text("原始误差 vs 真值", "Raw error vs reference"),
    )
    figure.colorbar(
        scatter_residual,
        ax=axes[0, 1],
        label=plot_text("频率 (kHz)", "Frequency (kHz)"),
    )

    axes[1, 0].scatter(frequency, percent, alpha=0.72, s=25)
    axes[1, 0].axhline(0, color="r", linestyle="--", linewidth=1.2)
    axes[1, 0].set(
        xlabel=plot_text("频率 (kHz)", "Frequency (kHz)"),
        ylabel=plot_text("原始相对误差 (%)", "Raw relative error (%)"),
        title=plot_text("频率分组误差", "Error by frequency"),
    )

    gain = np.divide(
        truth,
        measured,
        out=np.full_like(truth, np.nan),
        where=np.abs(measured) > EPS,
    )
    gain_frame = pd.DataFrame({"频率": frequency, "增益": gain}).dropna()
    grouped = gain_frame.groupby("频率")["增益"]
    gain_mean = grouped.mean()
    gain_std = grouped.std().fillna(0.0)
    axes[1, 1].errorbar(
        gain_mean.index,
        gain_mean.values,
        yerr=gain_std.values,
        marker="o",
        capsize=4,
        linewidth=1.4,
    )
    axes[1, 1].set(
        xlabel=plot_text("频率 (kHz)", "Frequency (kHz)"),
        ylabel=plot_text("真值 / 测量值", "Reference / measured"),
        title=plot_text("各频率所需校准增益（均值±标准差）", "Gain by frequency"),
    )
    axes[1, 1].grid(alpha=0.25)
    figure.suptitle(plot_text("校准数据概览", "Calibration data overview"), fontsize=15)
    figure.tight_layout()
    return save_figure(figure, output_dir, "01_数据概览", dpi, legacy_stem="原始残差图")


def plot_model_comparison(
    comparison: pd.DataFrame, output_dir: Path, dpi: int
) -> list[Path]:
    ordered = comparison.iloc[::-1].reset_index(drop=True)
    names = ordered["模型"].tolist()
    fit_rmse = ordered["拟合RMSE_mV"].to_numpy(dtype=float)
    cv_rmse = ordered["折RMSE均值_mV"].to_numpy(dtype=float)
    cv_se = ordered["折RMSE标准误_mV"].fillna(0).to_numpy(dtype=float)
    has_cv = np.isfinite(cv_rmse).any()

    figure, axes = plt.subplots(1, 2, figsize=(14, max(6, len(names) * 0.62)))
    positions = np.arange(len(names))
    if has_cv:
        axes[0].errorbar(
            cv_rmse,
            positions,
            xerr=cv_se,
            fmt="o",
            color="#4c78a8",
            ecolor="#333333",
            capsize=3,
            linewidth=1.2,
        )
        for pos, value in zip(positions, cv_rmse):
            axes[0].text(value, pos, f" {value:.3f}", va="center", fontsize=8)
        axes[0].set_xscale("log")
        axes[0].set_xlabel(
            plot_text("折 RMSE 均值 ± 标准误 (mV)", "Fold RMSE mean ± SE (mV)")
        )
        axes[0].set_title(
            plot_text("交叉验证：越低越好", "Cross-validation: lower is better")
        )
    else:
        axes[0].barh(positions, fit_rmse, color="#4c78a8", alpha=0.85)
        axes[0].set_xlabel(plot_text("拟合 RMSE (mV)", "Fit RMSE (mV)"))
        axes[0].set_title(plot_text("整表拟合", "Full-data fit"))
    axes[0].set_yticks(positions, names)
    axes[0].grid(axis="x", alpha=0.22)

    calibrated = ordered.loc[ordered["模型标识"] != "identity"].reset_index(drop=True)
    calibrated_names = calibrated["模型"].tolist()
    calibrated_positions = np.arange(len(calibrated))
    calibrated_fit = calibrated["拟合RMSE_mV"].to_numpy(dtype=float)
    calibrated_cv = calibrated["折RMSE均值_mV"].to_numpy(dtype=float)
    widths = 0.38
    axes[1].barh(
        calibrated_positions - widths / 2,
        calibrated_fit,
        height=widths,
        label=plot_text("整表拟合", "Full fit"),
        color="#f2a541",
    )
    if has_cv:
        axes[1].barh(
            calibrated_positions + widths / 2,
            calibrated_cv,
            height=widths,
            label=plot_text("交叉验证", "Cross-validation"),
            color="#4c78a8",
        )
    axes[1].set_yticks(calibrated_positions, calibrated_names)
    axes[1].set_xlabel(plot_text("RMSE (mV)", "RMSE (mV)"))
    axes[1].set_title(
        plot_text("校准模型局部比较", "Calibrated models: fit vs validation")
    )
    baseline = ordered.loc[ordered["模型标识"] == "identity"]
    if not baseline.empty:
        baseline_value = float(
            baseline.iloc[0]["折RMSE均值_mV"]
            if has_cv
            else baseline.iloc[0]["拟合RMSE_mV"]
        )
        axes[1].text(
            0.98,
            0.02,
            plot_text(
                f"未校准基线：{baseline_value:.3f} mV",
                f"Raw baseline: {baseline_value:.3f} mV",
            ),
            transform=axes[1].transAxes,
            ha="right",
            va="bottom",
            bbox={"boxstyle": "round", "facecolor": "white", "alpha": 0.8},
        )
    axes[1].legend()
    axes[1].grid(axis="x", alpha=0.22)
    figure.suptitle(
        plot_text("候选模型比较", "Candidate model comparison"), fontsize=15
    )
    figure.tight_layout()
    return save_figure(figure, output_dir, "02_模型比较", dpi)


def plot_recommended_diagnostics(
    data: pd.DataFrame,
    result: dict[str, Any],
    output_dir: Path,
    dpi: int,
) -> list[Path]:
    truth = data["真值"].to_numpy(dtype=float)
    frequency = data["频率"].to_numpy(dtype=float)
    fitted = result["fitted_prediction"]
    validation = (
        result["oof_prediction"] if result["oof_prediction"] is not None else fitted
    )
    label_prefix = "CV" if result["oof_prediction"] is not None else "拟合"
    fit_residual = fitted - truth
    validation_residual = validation - truth
    nonzero = np.abs(truth) > EPS
    percent = np.full_like(truth, np.nan)
    percent[nonzero] = (
        100.0 * np.abs(validation_residual[nonzero]) / np.abs(truth[nonzero])
    )

    figure, axes = plt.subplots(2, 2, figsize=(13, 9))
    scatter = axes[0, 0].scatter(
        truth, fitted, c=frequency, cmap="viridis", alpha=0.78, s=28
    )
    low = min(float(np.min(truth)), float(np.min(fitted)))
    high = max(float(np.max(truth)), float(np.max(fitted)))
    axes[0, 0].plot([low, high], [low, high], "r--", linewidth=1.2)
    axes[0, 0].set(
        xlabel=plot_text("真值 (mV Vpp)", "Reference (mV Vpp)"),
        ylabel=plot_text("整表拟合预测 (mV Vpp)", "Full-fit prediction (mV Vpp)"),
        title=plot_text("整表拟合一致性", "Full-fit agreement"),
    )
    figure.colorbar(
        scatter,
        ax=axes[0, 0],
        label=plot_text("频率 (kHz)", "Frequency (kHz)"),
    )

    axes[0, 1].scatter(fitted, fit_residual, c=frequency, cmap="viridis", s=28)
    axes[0, 1].axhline(0, color="r", linestyle="--", linewidth=1.2)
    axes[0, 1].set(
        xlabel=plot_text("整表拟合预测 (mV)", "Full-fit prediction (mV)"),
        ylabel=plot_text("拟合残差 (mV)", "Fit residual (mV)"),
        title=plot_text("拟合残差结构", "Fit residual structure"),
    )

    validation_scatter = axes[1, 0].scatter(
        frequency,
        validation_residual,
        c=truth,
        cmap="plasma",
        alpha=0.78,
        s=28,
    )
    axes[1, 0].axhline(0, color="r", linestyle="--", linewidth=1.2)
    axes[1, 0].set(
        xlabel=plot_text("频率 (kHz)", "Frequency (kHz)"),
        ylabel=plot_text(f"{label_prefix}残差 (mV)", f"{label_prefix} residual (mV)"),
        title=plot_text(
            f"{label_prefix}残差 vs 频率",
            f"{label_prefix} residual by frequency",
        ),
    )
    figure.colorbar(
        validation_scatter,
        ax=axes[1, 0],
        label=plot_text("真值 (mV)", "Reference (mV)"),
    )

    finite_percent = percent[np.isfinite(percent)]
    axes[1, 1].hist(
        finite_percent,
        bins=min(24, max(8, int(np.sqrt(len(finite_percent))))),
        color="#4c78a8",
        edgecolor="white",
        alpha=0.88,
    )
    p95 = float(np.percentile(finite_percent, 95)) if len(finite_percent) else 0.0
    axes[1, 1].axvline(
        p95,
        color="r",
        linestyle="--",
        label=plot_text(f"P95={p95:.3f}%", f"P95={p95:.3f}%"),
    )
    axes[1, 1].set(
        xlabel=plot_text(
            f"{label_prefix}绝对百分误差 (%)",
            f"{label_prefix} absolute percentage error (%)",
        ),
        ylabel=plot_text("样本数", "Count"),
        title=plot_text("验证相对误差分布", "Validation relative error"),
    )
    axes[1, 1].legend()
    figure.suptitle(
        plot_text(
            f"推荐模型诊断：{result['spec'].name}",
            f"Recommended model: {result['spec'].name}",
        ),
        fontsize=15,
    )
    figure.tight_layout()
    return save_figure(
        figure, output_dir, "03_推荐模型诊断", dpi, legacy_stem="校准后残差图"
    )


def plot_frequency_performance(
    summary: pd.DataFrame,
    output_dir: Path,
    dpi: int,
    has_cross_validation: bool,
) -> list[Path]:
    frequency = summary["频率_kHz"].to_numpy(dtype=float)
    labels = [f"{value:g}" for value in frequency]
    positions = np.arange(len(frequency))
    width = 0.36

    figure, axes = plt.subplots(1, 2, figsize=(14, 5.8))
    axes[0].bar(
        positions - width / 2,
        summary["原始MAE_mV"],
        width,
        label=plot_text("校准前", "Raw"),
        color="#b8b8b8",
    )
    axes[0].bar(
        positions + width / 2,
        summary["CV_MAE_mV"],
        width,
        label=plot_text(
            "推荐模型 CV" if has_cross_validation else "推荐模型拟合",
            "Recommended CV" if has_cross_validation else "Recommended fit",
        ),
        color="#4c78a8",
    )
    axes[0].set_xticks(positions, labels)
    axes[0].set(
        xlabel=plot_text("频率 (kHz)", "Frequency (kHz)"),
        ylabel=plot_text("MAE (mV)", "MAE (mV)"),
        title=plot_text("各频率绝对误差", "Absolute error by frequency"),
    )
    axes[0].legend()
    axes[0].grid(axis="y", alpha=0.22)

    axes[1].axhline(0, color="black", linewidth=1)
    axes[1].plot(
        positions,
        summary["拟合偏差_mV"],
        "o-",
        label=plot_text("整表拟合", "Full fit"),
        color="#f2a541",
    )
    if has_cross_validation:
        axes[1].plot(
            positions,
            summary["CV偏差_mV"],
            "o-",
            label=plot_text("交叉验证", "Cross-validation"),
            color="#4c78a8",
        )
    axes[1].set_xticks(positions, labels)
    axes[1].set(
        xlabel=plot_text("频率 (kHz)", "Frequency (kHz)"),
        ylabel=plot_text("平均偏差 (mV)", "Mean bias (mV)"),
        title=plot_text("校准后各频率系统偏差", "Post-calibration bias by frequency"),
    )
    axes[1].legend()
    axes[1].grid(axis="y", alpha=0.22)
    figure.tight_layout()
    return save_figure(figure, output_dir, "04_分频率表现", dpi)


# --------------------------- 控制台与文件报告 ---------------------------
def print_console_summary(
    data: pd.DataFrame,
    audit: dict[str, Any],
    comparison: pd.DataFrame,
    results: dict[str, dict[str, Any]],
    recommended_key: str,
    selection: dict[str, Any],
    cv_description: str,
    formula: dict[str, str],
    ranges: dict[str, float],
) -> None:
    console.print(
        Panel.fit(
            "\n".join(
                [
                    f"样本：{len(data)}，频率组：{data['频率'].nunique()}",
                    (
                        f"测量范围：{ranges['voltage_min_mv']:.3f}–"
                        f"{ranges['voltage_max_mv']:.3f} mV Vpp"
                    ),
                    (
                        f"频率范围：{ranges['frequency_min_khz']:.3f}–"
                        f"{ranges['frequency_max_khz']:.3f} kHz"
                    ),
                    f"验证策略：{cv_description}",
                    f"输入 SHA-256：{audit['input_sha256'][:16]}…",
                ]
            ),
            title="数据摘要",
        )
    )

    table = Table(title="候选模型：拟合与泛化能力", show_lines=False)
    table.add_column("推荐", justify="center")
    table.add_column("模型")
    table.add_column("参数", justify="right")
    table.add_column("拟合 RMSE", justify="right")
    table.add_column("折均 RMSE", justify="right")
    table.add_column("OOF MAE", justify="right")
    table.add_column("改善", justify="right")
    for _, row in comparison.iterrows():
        cv_rmse = row["折RMSE均值_mV"]
        table.add_row(
            "[green]●[/green]" if row["推荐"] == "是" else "",
            str(row["模型"]),
            str(int(row["参数数量"])),
            format_number(row["拟合RMSE_mV"]),
            format_number(cv_rmse),
            format_number(row["CV_MAE_mV"]),
            (
                f"{row['相对未校准RMSE改善_percent']:.1f}%"
                if pd.notna(row["相对未校准RMSE改善_percent"])
                else "—"
            ),
        )
    console.print(table)

    recommended = results[recommended_key]
    fit = recommended["fit_metrics"]
    cv = recommended["cv_metrics"] or fit
    fold = recommended["fold_summary"]
    if fold:
        evaluation_line = (
            f"交叉验证：折均RMSE={fold['rmse_mean_mv']:.4f}±"
            f"{fold['rmse_se_mv']:.4f} mV；OOF汇总"
            f"RMSE={cv['rmse_mv']:.4f} mV，MAE={cv['mae_mv']:.4f} mV，"
            f"P95={cv['p95_abs_error_mv']:.4f} mV，"
            f"最大误差={cv['max_abs_error_mv']:.4f} mV"
        )
    else:
        evaluation_line = (
            "未执行交叉验证；以上及以下误差均为拟合内指标，"
            f"P95={fit['p95_abs_error_mv']:.4f} mV，"
            f"最大误差={fit['max_abs_error_mv']:.4f} mV"
        )
    definition = formula["definitions"]
    if len(definition) > 125:
        definition = definition[:122] + "…"
    console.print(
        Panel(
            "\n".join(
                [
                    f"模型：{recommended['spec'].name}",
                    f"选择规则：{selection['criterion']}",
                    f"整表拟合：MAE={fit['mae_mv']:.4f} mV，"
                    f"RMSE={fit['rmse_mv']:.4f} mV，R2={format_r2(fit['r2'])}",
                    evaluation_line,
                    f"公式：{formula['vpp']}",
                    f"定义：{definition}",
                    "提示：有效范围外属于未经验证的外推。",
                ]
            ),
            title="推荐结果",
            border_style="green",
        )
    )


def dataframe_markdown(
    frame: pd.DataFrame,
    columns: Sequence[str],
    float_digits: int = 4,
) -> str:
    header = "| " + " | ".join(columns) + " |"
    separator = "|" + "|".join(["---"] * len(columns)) + "|"
    rows = [header, separator]
    for row_index in frame.index:
        values = []
        for column in columns:
            value = frame.at[row_index, column]
            if pd.api.types.is_integer_dtype(frame[column].dtype):
                values.append(str(int(value)))
            elif isinstance(value, (float, np.floating)):
                values.append(format_number(value, float_digits))
            else:
                values.append(str(value))
        rows.append("| " + " | ".join(values) + " |")
    return "\n".join(rows)


def build_markdown_report(
    input_path: Path,
    audit: dict[str, Any],
    data: pd.DataFrame,
    comparison: pd.DataFrame,
    frequency_summary: pd.DataFrame,
    result: dict[str, Any],
    selection: dict[str, Any],
    cv_description: str,
    formula: dict[str, str],
    normalization: Normalization,
    ranges: dict[str, float],
    monotonicity: dict[str, Any],
    outlier_info: dict[str, Any],
    skipped: dict[str, str],
    generated_at: str,
) -> str:
    fit = result["fit_metrics"]
    has_cv = result["cv_metrics"] is not None
    validation = result["cv_metrics"] or fit
    evaluation_label = "交叉验证 OOF 汇总" if has_cv else "拟合内评估"
    evaluation_short = "验证" if has_cv else "拟合"
    baseline_row = comparison.loc[comparison["模型标识"] == "identity"].iloc[0]
    baseline_cv_rmse = baseline_row["CV_RMSE_mV"]
    baseline_rmse = (
        float(baseline_cv_rmse)
        if pd.notna(baseline_cv_rmse)
        else float(baseline_row["拟合RMSE_mV"])
    )
    improvement = safe_percent(
        baseline_rmse - validation["rmse_mv"],
        baseline_rmse,
    )
    model_table = comparison[
        [
            "推荐",
            "模型",
            "参数数量",
            "拟合RMSE_mV",
            "折RMSE均值_mV",
            "CV_MAE_mV",
            "CV_P95绝对误差_mV",
            "CV_R2",
        ]
    ].copy()
    model_table.columns = [
        "推荐",
        "模型",
        "参数数",
        "拟合RMSE(mV)",
        "折RMSE均值(mV)",
        "CV MAE(mV)",
        "CV P95(mV)",
        "CV R²",
    ]
    frequency_table = frequency_summary[
        [
            "频率_kHz",
            "样本数",
            "原始MAE_mV",
            "拟合MAE_mV",
            "CV_MAE_mV",
            "CV偏差_mV",
            "CV最大绝对误差_mV",
        ]
    ].copy()
    frequency_table.columns = [
        "频率(kHz)",
        "样本数",
        "原始MAE(mV)",
        "拟合MAE(mV)",
        "CV MAE(mV)",
        "CV偏差(mV)",
        "CV最大误差(mV)",
    ]
    coefficients = coefficient_dataframe(result)
    skip_text = (
        "\n".join(f"- {name}：{reason}" for name, reason in skipped.items())
        if skipped
        else "无。"
    )
    monotonic_text = (
        "通过"
        if monotonicity["monotonic_on_calibration_domain"]
        else "未通过（不可把该模型直接解释为单调瞬时值校准）"
    )
    suspects = outlier_info["local_consistency"]["suspects"]
    suspect_text = (
        "\n".join(
            (
                f"- 源行 {item['source_row']}：f={item['frequency_khz']:g} kHz，"
                f"测量={item['measured_mv']:.6g} mV，"
                f"真值={item['reference_mv']:.6g} mV，"
                f"局部残差={item['local_residual_mv']:+.4f} mV，"
                f"稳健分数={item['robust_score']:.2f}。"
            )
            for item in suspects
        )
        if suspects
        else "未发现超过局部一致性阈值的记录。"
    )
    fold = result["fold_summary"]
    amplitude_validation = result["amplitude_cv_metrics"]
    fold_text = (
        f"各折 RMSE 均值 {fold['rmse_mean_mv']:.4f} ± "
        f"{fold['rmse_se_mv']:.4f} mV（均值 ± 标准误）"
        if fold
        else "未执行交叉验证"
    )
    selection_method_text = (
        "模型推荐使用一标准误差规则，因此性能差异不显著时优先选择参数更少的模型。"
        if has_cv
        else f"未执行交叉验证，模型按{selection['criterion']}推荐。"
    )
    final_validation_text = (
        "当前交叉验证同时参与模型选择，因此不是完全独立的最终验收。正式定型前"
        "应使用另一批、不同时间采集的数据作一次只评估不调参的确认试验。"
        if has_cv
        else "本次未执行交叉验证，拟合内指标不能估计泛化性能；正式定型前至少"
        "应增加独立验证数据。"
    )
    return f"""# Vpp 校准分析报告

生成时间：{generated_at}  
程序版本：{APP_VERSION}  
输入文件：`{input_path}`  
输入 SHA-256：`{audit["input_sha256"]}`

## 结论

推荐模型：**{result["spec"].name}**。<br>
选择方法：{selection["criterion"]}。{selection["reason"]}<br>
相对未校准基线，{evaluation_short} RMSE 改善
**{format_number(improvement, 2)}%**。

- 整表拟合：MAE {fit["mae_mv"]:.4f} mV，RMSE {fit["rmse_mv"]:.4f} mV，
  R² {format_r2(fit["r2"])}。
- {evaluation_label}：MAE {validation["mae_mv"]:.4f} mV，
  RMSE {validation["rmse_mv"]:.4f} mV；{fold_text}。
- {evaluation_short} P95 绝对误差
  {validation["p95_abs_error_mv"]:.4f} mV，最大绝对误差
  {validation["max_abs_error_mv"]:.4f} mV。
- {evaluation_short}点中 |误差|≤1 mV 的比例为
  {validation["within_1mv_percent"]:.2f}%，|误差|≤2 mV 的比例为
  {validation["within_2mv_percent"]:.2f}%。
- 按参考幅值整组留出验证：RMSE {amplitude_validation["rmse_mv"]:.4f} mV，
  P95 绝对误差 {amplitude_validation["p95_abs_error_mv"]:.4f} mV。该指标专门用于
  识别高阶幅值项对现有幅值等级的过拟合。
- 瞬时等效式在校准范围内的单调性检查：{monotonic_text}；
  数值网格最小导数为 {monotonicity["minimum_derivative"]:.6g}。

## 数据覆盖与质量

- 有效样本：{len(data)}；频率组：{data["频率"].nunique()}。
- 测量 Vpp 范围：{ranges["voltage_min_mv"]:.6g}–
  {ranges["voltage_max_mv"]:.6g} mV。
- 真值 Vpp 范围：{ranges["truth_min_mv"]:.6g}–
  {ranges["truth_max_mv"]:.6g} mV。
- 频率范围：{ranges["frequency_min_khz"]:.6g}–
  {ranges["frequency_max_khz"]:.6g} kHz。
- 删除的全空行：{audit["blank_rows_removed"]}；完全重复记录涉及
  {audit["exact_duplicate_rows"]} 行。
- CV 残差与同频局部一致性诊断共标记
  {int(outlier_info["flagged_count"])} 点；
  程序没有自动删除或重拟合这些点。

### 建议复核的记录

{suspect_text}

## 方法

旧方法为每条 Vpp 记录合成正弦，再以 Nelder–Mead 搜索参数。这并未增加观测
信息，还会形成高成本、非光滑的峰峰值目标。新版直接拟合 Vpp，并把模型写成
零截距的奇函数等效形式：

```text
{formula["definitions"]}
{formula["vpp"]}
{formula["instantaneous"]}
```

多项式候选先归一化，再以线性最小二乘一次求解；分频率插值候选在每个
频率结点估计过原点增益，再作线性插值。验证策略为
**{cv_description}**；此外固定执行按参考幅值留一组验证。{selection_method_text}系数区间使用
{result["coefficient_ci_method"]}。该区间反映回归系数稳定性，不是仪器计量
不确定度。

## 模型比较

“拟合”表示使用全部数据训练后再评估同一批数据；“CV”表示该点未参与对应折
训练，更适合估计泛化误差。

{dataframe_markdown(model_table, model_table.columns)}

## 推荐模型系数

{dataframe_markdown(coefficients, coefficients.columns, 8)}

## 分频率表现

{dataframe_markdown(frequency_table, frequency_table.columns)}

## 使用边界

1. 模型只在上述 Vpp 与频率覆盖范围内得到验证；范围外数值是外推。
2. 当前数据只有 Vpp、频率和参考 Vpp，不能证明逐采样点波形已被正确恢复。
   `{formula["note"]}`
3. 交叉验证衡量“未见过该频率组”时的表现；整表拟合指标仅反映对当前校准表
   的复现程度。
4. 若后续增加校准点，应重新运行本脚本并比较验证误差和分频率偏差，而不是
   只比较训练 R²。
5. {final_validation_text}

## 未能评估的候选模型

{skip_text}

## 输出文件

- `校准结果.xlsx`：推荐模型明细、全部模型预测、模型比较、分频率统计和系数。
- `模型比较.csv`、`校准明细.csv`、`分频率统计.csv`：便于程序读取。
- `推荐模型.json`：全精度系数、归一化参数、指标和有效范围。
- `校准模型.h`：可移植的 C 公式与范围检查。
- `01_数据概览` 至 `04_分频率表现`：PNG 与 SVG 图。
- `模型指标.txt`、`原始残差图.png`、`校准后残差图.png`：旧文件名兼容输出。
"""


def save_excel_workbook(
    path: Path,
    detail: pd.DataFrame,
    all_predictions: pd.DataFrame,
    comparison: pd.DataFrame,
    frequency_summary: pd.DataFrame,
    coefficients: pd.DataFrame,
    audit: dict[str, Any],
) -> None:
    audit_frame = pd.DataFrame(
        {
            "项目": [
                "输入文件",
                "输入SHA256",
                "工作表",
                "样本数",
                "频率组数",
                "删除全空行",
                "完全重复记录行数",
                "重复条件记录行数",
            ],
            "值": [
                audit["input_path"],
                audit["input_sha256"],
                audit["sheet"],
                audit["sample_count"],
                audit["frequency_count"],
                audit["blank_rows_removed"],
                audit["exact_duplicate_rows"],
                audit["repeated_condition_rows"],
            ],
        }
    )
    try:
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "写入 Excel 需要 openpyxl；请执行：pip install openpyxl"
        ) from exc

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        detail.to_excel(writer, sheet_name="推荐模型明细", index=False)
        all_predictions.to_excel(writer, sheet_name="全部模型预测", index=False)
        comparison.to_excel(writer, sheet_name="模型比较", index=False)
        frequency_summary.to_excel(writer, sheet_name="分频率统计", index=False)
        coefficients.to_excel(writer, sheet_name="推荐模型系数", index=False)
        audit_frame.to_excel(writer, sheet_name="数据审计", index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        recommended_fill = PatternFill("solid", fgColor="E2F0D9")
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in worksheet.columns:
                values = [str(cell.value or "") for cell in column_cells[:200]]
                width = min(max(max(map(len, values), default=8) + 2, 10), 30)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width

        comparison_sheet = workbook["模型比较"]
        headers = {cell.value: cell.column for cell in comparison_sheet[1]}
        if "推荐" in headers:
            recommend_column = headers["推荐"]
            for row in range(2, comparison_sheet.max_row + 1):
                if comparison_sheet.cell(row, recommend_column).value == "是":
                    for cell in comparison_sheet[row]:
                        cell.fill = recommended_fill

        detail_sheet = workbook["推荐模型明细"]
        detail_headers = {cell.value: cell.column_letter for cell in detail_sheet[1]}
        for header, column_letter in detail_headers.items():
            if isinstance(header, str) and "误差_mV" in header:
                detail_sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{detail_sheet.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="63BE7B",
                        end_type="max",
                        end_color="F8696B",
                    ),
                )


def save_outputs(
    output_dir: Path,
    input_path: Path,
    data: pd.DataFrame,
    audit: dict[str, Any],
    results: dict[str, dict[str, Any]],
    skipped: dict[str, str],
    recommended_key: str,
    selection: dict[str, Any],
    cv_description: str,
    normalization: Normalization,
    ranges: dict[str, float],
    comparison: pd.DataFrame,
    frequency_summary: pd.DataFrame,
    detail: pd.DataFrame,
    all_predictions: pd.DataFrame,
    outlier_info: dict[str, Any],
    formula: dict[str, str],
    monotonicity: dict[str, Any],
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    result = results[recommended_key]
    coefficients = coefficient_dataframe(result)
    written: list[Path] = []

    csv_frames = {
        "模型比较.csv": comparison,
        "校准明细.csv": detail,
        "分频率统计.csv": frequency_summary,
    }
    for filename, frame in csv_frames.items():
        path = output_dir / filename
        frame.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.12g")
        written.append(path)

    excel_path = output_dir / "校准结果.xlsx"
    save_excel_workbook(
        excel_path,
        detail,
        all_predictions,
        comparison,
        frequency_summary,
        coefficients,
        audit,
    )
    written.append(excel_path)

    model_payload = {
        "schema_version": 1,
        "application_version": APP_VERSION,
        "generated_at": generated_at,
        "source": audit,
        "units": {
            "measured_vpp": "mV",
            "reference_vpp": "mV",
            "frequency": "kHz",
        },
        "valid_range": ranges,
        "cv": {"description": cv_description},
        "selection": selection,
        "normalization": {
            "voltage_scale_mv": normalization.voltage_scale_mv,
            "frequency_center_khz": normalization.frequency_center_khz,
            "frequency_scale_khz": normalization.frequency_scale_khz,
        },
        "recommended_model": {
            "key": result["spec"].key,
            "name": result["spec"].name,
            "description": result["spec"].description,
            "parameter_count": result["parameter_count"],
            "formula": formula,
            "coefficients": result["coefficient_rows"],
            "coefficient_ci_method": result["coefficient_ci_method"],
            "design_matrix_condition_number": result["condition_number"],
            "monotonicity_check": monotonicity,
            "fit_metrics": result["fit_metrics"],
            "cross_validation_metrics": result["cv_metrics"],
            "amplitude_holdout_metrics": result["amplitude_cv_metrics"],
            "amplitude_holdout_folds": result["amplitude_fold_rows"],
            "fold_summary": result["fold_summary"],
            "folds": result["fold_rows"],
        },
        "outlier_diagnostic": outlier_info,
        "skipped_models": skipped,
    }
    json_path = output_dir / "推荐模型.json"
    with json_path.open("w", encoding="utf-8") as stream:
        json.dump(
            to_builtin(model_payload),
            stream,
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        )
    written.append(json_path)

    header_path = output_dir / "校准模型.h"
    header_path.write_text(
        build_c_header(
            result["spec"],
            result["coefficients"],
            normalization,
            ranges,
            audit["input_sha256"],
        ),
        encoding="utf-8",
    )
    written.append(header_path)

    report = build_markdown_report(
        input_path,
        audit,
        data,
        comparison,
        frequency_summary,
        result,
        selection,
        cv_description,
        formula,
        normalization,
        ranges,
        monotonicity,
        outlier_info,
        skipped,
        generated_at,
    )
    report_path = output_dir / "校准报告.md"
    report_path.write_text(report, encoding="utf-8")
    written.append(report_path)

    legacy_lines = [
        "Vpp 校准模型指标",
        "=" * 72,
        f"生成时间：{generated_at}",
        f"推荐模型：{result['spec'].name}",
        f"选择规则：{selection['criterion']}",
        f"验证策略：{cv_description}",
        f"公式：{formula['vpp']}",
        f"归一化：{formula['definitions']}",
        "",
        comparison.to_string(index=False),
        "",
        "完整说明请查看：校准报告.md",
    ]
    legacy_path = output_dir / "模型指标.txt"
    legacy_path.write_text("\n".join(legacy_lines), encoding="utf-8")
    written.append(legacy_path)
    return written


# --------------------------- 命令行与主流程 ---------------------------
def parse_sheet(value: str) -> str | int:
    return int(value) if value.isdigit() else value


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "根据测量 Vpp、参考 Vpp 和频率建立校准模型，使用交叉验证选模，"
            "并输出完整报告、图表和可部署系数。"
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "input",
        nargs="?",
        default="校准.xlsx",
        help="输入 Excel/CSV 文件",
    )
    parser.add_argument(
        "output",
        nargs="?",
        default="cb_output",
        help="输出目录；只覆盖本程序的同名结果，不清空目录",
    )
    parser.add_argument(
        "--sheet",
        default="0",
        help="Excel 工作表序号（从 0 开始）或名称",
    )
    parser.add_argument(
        "--cv",
        choices=["auto", "frequency", "kfold", "none"],
        default="auto",
        help="验证方式；auto 在有多个频率时按频率留组",
    )
    parser.add_argument(
        "--folds",
        type=int,
        default=5,
        help="随机 K 折模式的折数",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=RANDOM_SEED,
        help="随机 K 折种子",
    )
    parser.add_argument(
        "--no-plots",
        action="store_true",
        help="不生成 PNG/SVG 图，其他报告照常输出",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=200,
        help="PNG 分辨率",
    )
    parser.add_argument(
        "--plain",
        action="store_true",
        help="关闭终端颜色，便于日志或旧终端显示",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="出错时显示完整异常堆栈",
    )
    return parser


def resolve_input_path(raw_path: str) -> Path:
    path = Path(raw_path).expanduser()
    if path.is_file():
        return path.resolve()
    if not path.is_absolute() and raw_path == "校准.xlsx":
        repository_default = Path(__file__).resolve().parent.parent / raw_path
        if repository_default.is_file():
            return repository_default
    return path.resolve()


def run(args: argparse.Namespace) -> dict[str, Any]:
    input_path = resolve_input_path(args.input)
    sheet = parse_sheet(str(args.sheet))
    console.print(f"[bold]读取数据：[/bold]{input_path}")
    data, audit = load_data(input_path, sheet)

    measured = data["测量值"].to_numpy(dtype=float)
    truth = data["真值"].to_numpy(dtype=float)
    frequency = data["频率"].to_numpy(dtype=float)
    normalization = make_normalization(measured, frequency)
    splits, cv_description = build_cv_splits(frequency, args.cv, args.folds, args.seed)

    console.print(f"评估 {len(MODELS)} 个候选模型；{cv_description}…")
    results, skipped = evaluate_models(
        measured, frequency, truth, normalization, splits
    )
    recommended_key, selection = select_recommended_model(results)
    recommended = results[recommended_key]
    validation_prediction = (
        recommended["oof_prediction"]
        if recommended["oof_prediction"] is not None
        else recommended["fitted_prediction"]
    )

    raw_result = results.get("identity")
    if raw_result is None:
        raise RuntimeError("未校准基线评估失败，无法生成比较报告。")
    raw_rmse = (
        raw_result["cv_metrics"]["rmse_mv"]
        if raw_result["cv_metrics"] is not None
        else raw_result["fit_metrics"]["rmse_mv"]
    )
    comparison = model_comparison_dataframe(results, recommended_key, raw_rmse)
    frequency_summary = frequency_summary_dataframe(
        data,
        recommended["fitted_prediction"],
        validation_prediction,
    )
    detail, all_predictions, outlier_info = build_detail_dataframes(
        data, results, recommended_key
    )
    ranges = {
        "voltage_min_mv": float(np.min(measured)),
        "voltage_max_mv": float(np.max(measured)),
        "truth_min_mv": float(np.min(truth)),
        "truth_max_mv": float(np.max(truth)),
        "frequency_min_khz": float(np.min(frequency)),
        "frequency_max_khz": float(np.max(frequency)),
    }
    formula = model_formula(
        recommended["spec"],
        recommended["coefficients"],
        normalization,
    )
    monotonicity = check_instantaneous_monotonicity(
        recommended["spec"],
        recommended["coefficients"],
        normalization,
        ranges["voltage_max_mv"],
        ranges["frequency_min_khz"],
        ranges["frequency_max_khz"],
    )

    output_dir = Path(args.output).expanduser().resolve()
    written = save_outputs(
        output_dir,
        input_path,
        data,
        audit,
        results,
        skipped,
        recommended_key,
        selection,
        cv_description,
        normalization,
        ranges,
        comparison,
        frequency_summary,
        detail,
        all_predictions,
        outlier_info,
        formula,
        monotonicity,
    )

    if not args.no_plots:
        if not HAS_CHINESE_FONT:
            console.print(
                "[yellow]未找到中文绘图字体，图内文字将使用英文；数据文件仍为中文。[/yellow]"
            )
        written.extend(plot_data_overview(data, output_dir, args.dpi))
        written.extend(plot_model_comparison(comparison, output_dir, args.dpi))
        written.extend(
            plot_recommended_diagnostics(data, recommended, output_dir, args.dpi)
        )
        written.extend(
            plot_frequency_performance(
                frequency_summary,
                output_dir,
                args.dpi,
                recommended["oof_prediction"] is not None,
            )
        )

    print_console_summary(
        data,
        audit,
        comparison,
        results,
        recommended_key,
        selection,
        cv_description,
        formula,
        ranges,
    )
    suspects = outlier_info["local_consistency"]["suspects"]
    if suspects:
        console.print("[yellow]发现建议复核的局部不一致记录（未自动删除）：[/yellow]")
        for item in suspects:
            console.print(
                f"  - 源行 {item['source_row']}，"
                f"{item['frequency_khz']:g} kHz，"
                f"测量 {item['measured_mv']:.6g} mV / "
                f"真值 {item['reference_mv']:.6g} mV，"
                f"局部残差 {item['local_residual_mv']:+.4f} mV"
            )
    if audit["exact_duplicate_rows"]:
        console.print(
            f"[yellow]数据中有 {audit['exact_duplicate_rows']} 行完全重复记录；"
            "已保留并按实际权重参与拟合。[/yellow]"
        )
    if skipped:
        console.print("[yellow]部分模型因数据不足或矩阵秩不足而跳过：[/yellow]")
        for name, reason in skipped.items():
            console.print(f"  - {name}：{reason}")
    console.print(
        f"[bold green]完成。[/bold green]已写入 {len(written)} 个文件：{output_dir}"
    )
    console.print(f"建议先查看：{output_dir / '校准报告.md'}")
    return {
        "recommended_key": recommended_key,
        "recommended_name": recommended["spec"].name,
        "output_dir": str(output_dir),
        "written_files": [str(path) for path in written],
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_argument_parser()
    args = parser.parse_args(argv)
    global console
    console = Console(highlight=False, no_color=args.plain)
    if args.folds < 2:
        parser.error("--folds 必须至少为 2")
    if not 72 <= args.dpi <= 1200:
        parser.error("--dpi 必须在 72 到 1200 之间")
    try:
        run(args)
        return 0
    except KeyboardInterrupt:
        console.print("\n[yellow]用户中止。[/yellow]")
        return 130
    except Exception as exc:
        console.print(f"[bold red]错误：[/bold red]{exc}")
        if args.debug:
            raise
        console.print("可加 --debug 查看完整堆栈；用 -h 查看参数说明。")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
